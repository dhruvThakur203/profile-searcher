from pathlib import Path

from openpyxl import Workbook, load_workbook
from openpyxl.styles import Alignment

from ..models import PersonRecord
from .ai_overview import folio_from_headers_and_values
from .researcher import (
    PersonResearch,
    ResearchStatus,
    ShareholderResearcher,
    format_joint_holders,
    format_similarity,
    research_status_summary,
)
from .search_client import SearchClient, SearchPolicy


CORE_OUTPUT_COLUMNS = [
    "Matched Name",
    "Mobile",
    "Email",
    "Source Links",
    "Evidence",
    "Match Status",
    "Address 1 Match %",
    "Address 2 Match %",
    "Address 3 Match %",
    "Address 4 Match %",
    "Search Strategies Used",
]
# Appended after the core columns so existing positions/names never change.
EXTENDED_OUTPUT_COLUMNS = [
    "Match Basis",
    "Name Match Classification",
    "Joint Holder Evidence",
    "Obituary / Condolence Lead",
    "Family/Relationship Evidence",
    "Location Evidence",
    "Company Evidence",
    "Manual Review Reason",
    "Manual Review Sources",
    "Source Evidence",
    "Derived Postal Fragments",
    "Research Status / Errors",
    "Planned Strategies Not Executed",
]
OUTPUT_COLUMNS = CORE_OUTPUT_COLUMNS + EXTENDED_OUTPUT_COLUMNS
_WRAPPED_COLUMNS = (
    "Source Links",
    "Evidence",
    "Joint Holder Evidence",
    "Obituary / Condolence Lead",
    "Family/Relationship Evidence",
    "Manual Review Reason",
    "Manual Review Sources",
    "Source Evidence",
    "Research Status / Errors",
    "Planned Strategies Not Executed",
)
SOURCE_EVIDENCE_COLUMNS = [
    "Record Id",
    "Input Name",
    "Classification",
    "Match Basis",
    "Name Match Classification",
    "Address 1 %",
    "Address 2 %",
    "Address 3 %",
    "Address 4 %",
    "Address Fields >= 90%",
    "Joint Holder Evidence",
    "Record Group / Span",
    "Page",
    "Original URL",
    "Final URL",
    "Title",
    "Document Type",
    "Extraction Status",
    "Strategies",
    "Manual Review Reason",
    "Record Context",
]
AI_OVERVIEW_LOG_COLUMNS = [
    "Input Excel Row",
    "Folio Number",
    "Primary Holder",
    "Holder Associated With Query",
    "Query Strategy",
    "Exact Search Query",
    "Capture Timestamp",
    "Overview Status",
    "Raw Overview Text",
    "Source Titles",
    "Source URLs",
    "Extraction Notes",
]
AI_CONTACT_LEAD_COLUMNS = [
    "Person Name",
    "Email Address",
    "Phone Number",
    "Institution or Company",
    "Job Title or Role",
    "Address or Location",
    "Contact Type",
    "Input Excel Row",
    "Folio Number",
    "Holder Associated With Query",
    "Exact Search Query",
    "Supporting Source Title",
    "Supporting Source URL",
    "Supporting Evidence Text",
    "Verification Status",
    "Extraction Notes",
]
OBITUARY_LEAD_COLUMNS = [
    "Record Id",
    "Input Name",
    "Source URL",
    "Original URL",
    "Title",
    "Document Type",
    "Page",
    "Holder Name Mention",
    "Mentioned As",
    "Joint Holder Mentions",
    "Keywords",
    "Names Mentioned",
    "Relationship Wording",
    "Locations",
    "Companies (lead only)",
    "Excerpt",
    "Suggested Follow-up Queries (not executed)",
    "Found By Strategies",
    "Note",
]
MOBILE_SEPARATOR = " | "
EMAIL_SEPARATOR = " | "
SOURCE_LINK_SEPARATOR = "\n"
STRATEGY_SEPARATOR = " | "

_HEADER_FIELDS = {
    "name": "name",
    "shareholder name": "name",
    "person name": "name",
    "full name": "name",
    "name of the holder": "name",
    "holder name": "name",
    "second holder joint 1": "joint1",
    "second holder": "joint1",
    "third holder joint 2": "joint2",
    "third holder": "joint2",
    "address 1": "address1",
    "address1": "address1",
    "address line 1": "address1",
    "addr 1": "address1",
    "address 2": "address2",
    "address2": "address2",
    "address line 2": "address2",
    "addr 2": "address2",
    "address 3": "address3",
    "address3": "address3",
    "address line 3": "address3",
    "address 4": "address4",
    "address4": "address4",
    "address line 4": "address4",
    "pincode": "pincode",
    "pin code": "pincode",
    "postal code": "pincode",
    "zip": "pincode",
    "zipcode": "pincode",
    "country": "country",
    "scrip": "scrip",
    "company": "scrip",
    "scrip company": "scrip",
    "record id": "record_id",
    "id": "record_id",
}


class InputRow:
    def __init__(self, record: PersonRecord, values: list[object], excel_row: int) -> None:
        self.record = record
        self.values = values
        self.excel_row = excel_row


def run_workbook(
    input_path: str | Path,
    output_path: str | Path | None,
    client: SearchClient,
    policy: SearchPolicy | None = None,
    checkpoint=None,
    excel_row: int | None = None,
    excel_rows: list[int] | None = None,
    source_fetcher=None,
    source_inspection_top_n: int | None = None,
) -> list[PersonResearch]:
    source = Path(input_path).resolve()
    destination = None if output_path is None else Path(output_path).resolve()
    if not source.is_file():
        raise ValueError(
            f"Input workbook not found: {source}\n"
            "Pass the full path to your shareholder .xlsx file "
            "(for example input.xlsx in this project folder), not a placeholder like path\\to\\input.xlsx."
        )
    if destination is not None and source == destination:
        raise ValueError("Output path must be different from the input workbook.")

    sheet_name, headers, rows = load_input(source)
    selected = None
    if excel_rows is not None:
        wanted = set(excel_rows)
        selected = [row for row in rows if row.excel_row in wanted]
        missing = sorted(wanted - {row.excel_row for row in selected})
        if missing:
            raise ValueError(
                "No shareholder row at Excel row(s): "
                + ", ".join(str(value) for value in missing)
            )
        # Preserve the caller order for multi-row test runs.
        order = {value: index for index, value in enumerate(excel_rows)}
        selected.sort(key=lambda row: order.get(row.excel_row, row.excel_row))
        rows = selected
    elif excel_row is not None:
        rows = [row for row in rows if row.excel_row == excel_row]
        if not rows:
            raise ValueError(f"No shareholder row at Excel row {excel_row}.")
    researcher = ShareholderResearcher(
        client,
        policy,
        checkpoint=checkpoint,
        source_fetcher=source_fetcher,
        source_inspection_top_n=source_inspection_top_n,
    )
    researched = []
    interrupted = False
    try:
        for row in rows:
            researched.append((row, researcher.research_person(row.record)))
    except KeyboardInterrupt:
        interrupted = True
        # The row in progress and every row not reached are reported with
        # their unexecuted searches instead of being silently dropped.
        for row in rows[len(researched):]:
            researched.append((row, researcher.interrupted_outcome(row.record)))
    if destination is not None:
        write_output(destination, sheet_name, headers, researched)
    if interrupted:
        raise KeyboardInterrupt
    return [outcome for _, outcome in researched]


def load_input(path: str | Path) -> tuple[str, list[str], list[InputRow]]:
    workbook = load_workbook(Path(path), data_only=True)
    try:
        sheet = workbook.active
        raw_rows = list(sheet.iter_rows(values_only=True))
        sheet_name = sheet.title
    finally:
        workbook.close()

    if not raw_rows:
        raise ValueError("The workbook has no header row.")

    headers = [_header_text(value, index) for index, value in enumerate(raw_rows[0], start=1)]
    fields = [_HEADER_FIELDS.get(_normalize_header(header), "") for header in headers]
    if "name" not in fields:
        raise ValueError("The workbook must include a Name column.")

    rows = []
    for excel_row, raw in enumerate(raw_rows[1:], start=2):
        values = list(raw) + [None] * max(0, len(headers) - len(raw))
        values = values[: len(headers)]
        if all(_blank(value) for value in values):
            continue
        record = _record_from_row(excel_row, fields, values)
        rows.append(InputRow(record=record, values=values, excel_row=excel_row))
    return sheet_name, headers, rows


def write_output(path, sheet_name, headers, researched) -> None:
    workbook = Workbook()
    sheet = workbook.active
    sheet.title = sheet_name
    output_headers, positions = _output_header_plan(headers)
    sheet.append(output_headers)
    for row, outcome in researched:
        values = list(row.values) + [None] * max(0, len(output_headers) - len(row.values))
        values = values[: len(output_headers)]
        mapped = _output_map(outcome)
        for column, index in positions.items():
            values[index] = mapped[column]
        sheet.append(values)

    if sheet.max_row >= 2:
        wrap = Alignment(wrap_text=True, vertical="top")
        for column_name in _WRAPPED_COLUMNS:
            column = positions[column_name] + 1
            for row_cells in sheet.iter_rows(
                min_row=2,
                max_row=sheet.max_row,
                min_col=column,
                max_col=column,
            ):
                for cell in row_cells:
                    cell.alignment = wrap
    sheet.freeze_panes = "A2"

    log_title = "Research Log"
    if log_title == sheet_name:
        log_title = "Search History"
    log = workbook.create_sheet(log_title)
    log.append(
        [
            "Record Id",
            "Input Name",
            "Strategy",
            "Query",
            "Result Count",
            "Accepted Candidate Count",
            "Status",
            "Message",
            "Timestamp",
            "Investigation Notes",
            "Google Results",
        ]
    )
    for row, outcome in researched:
        if outcome.log:
            for entry in outcome.log:
                log.append(
                    [
                        outcome.record_id,
                        outcome.input_name,
                        entry.strategy,
                        entry.query,
                        entry.result_count,
                        entry.accepted_candidate_count,
                        entry.status.value,
                        entry.message,
                        entry.timestamp,
                        "\n".join(entry.investigation_notes or []),
                        _format_result_details(entry.result_details),
                    ]
                )
        else:
            log.append(
                [
                    outcome.record_id,
                    outcome.input_name,
                    "",
                    "",
                    0,
                    0,
                    outcome.status.value,
                    outcome.message,
                    "",
                    "",
                    "",
                ]
            )
        for item in _unexecuted_log_items(outcome):
            log.append(
                [
                    outcome.record_id,
                    outcome.input_name,
                    item["strategy"],
                    item["query"],
                    None,
                    None,
                    NOT_EXECUTED_STATUS,
                    item["reason"],
                    "",
                    "",
                    "",
                ]
            )

    taken = {sheet_name, log_title}
    evidence_sheet = workbook.create_sheet(_free_title("Source Evidence", taken))
    evidence_sheet.append(SOURCE_EVIDENCE_COLUMNS)
    for _, outcome in researched:
        for finding in list(outcome.strong_findings) + list(outcome.review_findings):
            evidence_sheet.append(_source_evidence_row(outcome, finding))
    evidence_sheet.freeze_panes = "A2"

    leads_sheet = workbook.create_sheet(_free_title("Obituary Leads", taken))
    leads_sheet.append(OBITUARY_LEAD_COLUMNS)
    for _, outcome in researched:
        for lead in outcome.obituary_leads:
            leads_sheet.append(_obituary_lead_row(outcome, lead))
    leads_sheet.freeze_panes = "A2"

    overview_sheet = workbook.create_sheet(_free_title("AI Overview Log", taken))
    overview_sheet.append(AI_OVERVIEW_LOG_COLUMNS)
    for row, outcome in researched:
        folio = folio_from_headers_and_values(headers, row.values)
        for entry in outcome.ai_overview_log:
            overview_sheet.append(_ai_overview_log_row(row.excel_row, folio, outcome, entry))
    overview_sheet.freeze_panes = "A2"

    ai_leads_sheet = workbook.create_sheet(_free_title("AI Contact Leads", taken))
    ai_leads_sheet.append(AI_CONTACT_LEAD_COLUMNS)
    for row, outcome in researched:
        folio = folio_from_headers_and_values(headers, row.values)
        for lead in outcome.ai_contact_leads:
            ai_leads_sheet.append(_ai_contact_lead_row(row.excel_row, folio, lead))
    ai_leads_sheet.freeze_panes = "A2"
    workbook.save(path)
    workbook.close()


NOT_EXECUTED_STATUS = "NOT_EXECUTED"


def _unexecuted_log_items(outcome) -> list[dict]:
    if outcome.unexecuted_plans:
        return outcome.unexecuted_plans
    # Outcomes saved before queries were recorded only have "STRATEGY (not executed: reason)".
    items = []
    for text in outcome.skipped_strategies:
        strategy, _, rest = text.partition(" (not executed: ")
        items.append({"strategy": strategy, "query": "", "reason": rest.rstrip(")") or text})
    return items


def _format_result_details(details: list[dict]) -> str:
    lines = []
    for item in details or []:
        line = f"#{item.get('rank', '')} {item.get('url', '')}"
        if item.get("title"):
            line += f" | {item['title']}"
        if item.get("snippet"):
            line += f" | {item['snippet']}"
        lines.append(line)
    # Excel cells hold at most 32767 characters.
    return "\n".join(lines)[:32000]


def _free_title(title: str, taken: set[str]) -> str:
    candidate = title
    index = 2
    while candidate in taken:
        candidate = f"{title} {index}"
        index += 1
    taken.add(candidate)
    return candidate


def _source_evidence_row(outcome: PersonResearch, finding) -> list[object]:
    context = ""
    for line in (finding.evidence_summary or "").splitlines():
        if line.startswith("Record context:"):
            context = line.removeprefix("Record context:").strip()
            break
    return [
        outcome.record_id,
        outcome.input_name,
        finding.kind,
        finding.match_basis,
        finding.name_equivalence,
        finding.address1_similarity,
        finding.address2_similarity,
        finding.address3_similarity,
        finding.address4_similarity,
        ", ".join(finding.matched_address_fields),
        format_joint_holders(finding.joint_holders),
        finding.evidence_span or finding.group_id,
        finding.page_number,
        finding.original_url or finding.url,
        finding.url,
        finding.title,
        finding.document_type,
        "OK",
        ", ".join(finding.strategies),
        finding.reason,
        context,
    ]


def _ai_overview_log_row(
    excel_row: int,
    folio: str,
    outcome: PersonResearch,
    entry,
) -> list[object]:
    return [
        excel_row,
        folio or entry.folio_number,
        outcome.input_name,
        entry.holder_associated,
        entry.query_strategy,
        entry.search_query,
        entry.capture_timestamp,
        entry.overview_status,
        entry.raw_overview_text,
        "\n".join(entry.source_titles),
        "\n".join(entry.source_urls),
        entry.extraction_notes,
    ]


def _ai_contact_lead_row(excel_row: int, folio: str, lead) -> list[object]:
    return [
        lead.person_name,
        lead.email,
        lead.phone,
        lead.institution,
        lead.job_title,
        lead.location,
        lead.contact_type,
        excel_row,
        folio or lead.folio_number,
        lead.holder_associated,
        lead.search_query,
        lead.supporting_source_title,
        lead.supporting_source_url,
        lead.supporting_evidence_text,
        lead.verification_status,
        lead.extraction_notes,
    ]


def _obituary_lead_row(outcome: PersonResearch, lead) -> list[object]:
    return [
        outcome.record_id,
        outcome.input_name,
        lead.source_url,
        lead.original_url or lead.source_url,
        lead.title,
        lead.document_type,
        lead.page_number,
        lead.primary_name_mention,
        lead.primary_name_text,
        "; ".join(f"{key}: {value}" for key, value in lead.joint_holder_mentions.items()),
        ", ".join(lead.keywords),
        " | ".join(lead.names),
        "\n".join(_relationship_text(item) for item in lead.relationships),
        " | ".join(lead.locations),
        " | ".join(lead.companies),
        lead.excerpt,
        "\n".join(lead.follow_up_queries),
        ", ".join(lead.strategies),
        lead.note,
    ]


def _relationship_text(item) -> str:
    text = f"{item.term} ({item.category}): \"{item.phrase}\""
    details = []
    if item.named:
        details.append("named: " + ", ".join(item.named))
    if item.preceding_names:
        details.append("preceding: " + ", ".join(item.preceding_names))
    if item.sentence_names:
        details.append("also in sentence: " + ", ".join(item.sentence_names))
    if details:
        text += " [" + "; ".join(details) + "]"
    return text


def _output_header_plan(headers: list[str]) -> tuple[list[str], dict[str, int]]:
    positions = {}
    normalized = {_normalize_header(header): index for index, header in enumerate(headers)}
    output_headers = list(headers)
    for column in OUTPUT_COLUMNS:
        key = _normalize_header(column)
        if key in normalized:
            positions[column] = normalized[key]
        else:
            positions[column] = len(output_headers)
            output_headers.append(column)
    return output_headers, positions


def _output_map(outcome: PersonResearch) -> dict[str, object]:
    return {
        "Matched Name": outcome.matched_name,
        "Mobile": MOBILE_SEPARATOR.join(outcome.mobiles),
        "Email": EMAIL_SEPARATOR.join(outcome.emails),
        "Source Links": SOURCE_LINK_SEPARATOR.join(outcome.source_links),
        "Evidence": outcome.evidence,
        "Match Status": outcome.status.value,
        "Address 1 Match %": outcome.address1_match_pct,
        "Address 2 Match %": outcome.address2_match_pct,
        "Address 3 Match %": outcome.address3_match_pct,
        "Address 4 Match %": outcome.address4_match_pct,
        "Search Strategies Used": STRATEGY_SEPARATOR.join(outcome.strategies_used),
        **_extended_output_map(outcome),
    }


def _extended_output_map(outcome: PersonResearch) -> dict[str, object]:
    scored = outcome.strong_findings or (
        outcome.review_findings if outcome.status == ResearchStatus.MANUAL_REVIEW else []
    )
    findings = list(outcome.strong_findings) + list(outcome.review_findings)
    leads = outcome.obituary_leads
    return {
        "Match Basis": " | ".join(_unique(item.match_basis for item in scored)),
        "Name Match Classification": " | ".join(
            _unique(item.name_equivalence for item in scored)
        ),
        "Joint Holder Evidence": "\n".join(
            f"[{item.kind}] {item.url}: {format_joint_holders(item.joint_holders)}"
            for item in findings
            if item.joint_holders
        ),
        "Obituary / Condolence Lead": "\n".join(
            f"{lead.source_url} | {lead.title} | holder mentioned as "
            f"{lead.primary_name_mention} ({lead.primary_name_text}) | {_clip(lead.excerpt)}"
            for lead in leads
        ),
        "Family/Relationship Evidence": "\n".join(
            f"{_relationship_text(item)} — {lead.source_url}"
            for lead in leads
            for item in lead.relationships
        ),
        "Location Evidence": "\n".join(
            f"{' | '.join(lead.locations)} — {lead.source_url}"
            for lead in leads
            if lead.locations
        ),
        "Company Evidence": "\n".join(
            f"{' | '.join(lead.companies)} (mentioned; relationship not established) "
            f"— {lead.source_url}"
            for lead in leads
            if lead.companies
        ),
        "Manual Review Reason": outcome.manual_review_reason,
        "Manual Review Sources": SOURCE_LINK_SEPARATOR.join(
            item.url for item in outcome.review_findings
        ),
        "Source Evidence": "\n".join(_source_evidence_line(item) for item in findings),
        "Derived Postal Fragments": "\n".join(outcome.postal_fragments),
        "Research Status / Errors": research_status_summary(outcome),
        "Planned Strategies Not Executed": "\n".join(outcome.skipped_strategies),
    }


def _source_evidence_line(finding) -> str:
    parts = [f"[{finding.kind}] {finding.url}"]
    if finding.original_url and finding.original_url != finding.url:
        parts.append(f"from {finding.original_url}")
    if finding.title:
        parts.append(finding.title)
    if finding.document_type:
        parts.append(finding.document_type)
    if finding.evidence_span:
        parts.append(finding.evidence_span)
    parts.append(f"name {finding.name_equivalence or 'EXACT'}")
    scores = [
        f"A{index} {format_similarity(value)}%"
        for index, value in enumerate(
            (
                finding.address1_similarity,
                finding.address2_similarity,
                finding.address3_similarity,
                finding.address4_similarity,
            ),
            start=1,
        )
        if value is not None
    ]
    if scores:
        parts.append(" ".join(scores))
    if finding.match_basis:
        parts.append(finding.match_basis)
    return " | ".join(parts)


def _clip(text: str, limit: int = 300) -> str:
    compact = " ".join((text or "").split())
    return compact if len(compact) <= limit else compact[: limit - 3] + "..."


def _unique(values) -> list[str]:
    seen = []
    for value in values:
        if value and value not in seen:
            seen.append(value)
    return seen


def _record_from_row(excel_row, fields, values) -> PersonRecord:
    data = {
        field: ""
        for field in (
            "name",
            "address1",
            "address2",
            "address3",
            "address4",
            "pincode",
            "country",
            "scrip",
            "joint1",
            "joint2",
        )
    }
    record_id = str(excel_row)
    for field, value in zip(fields, values):
        if not field:
            continue
        text = "" if value is None else str(value).strip()
        if field == "pincode" and text == "0":
            text = ""
        if field == "record_id":
            if text:
                record_id = text
            continue
        if field not in data or data[field]:
            continue
        data[field] = text
    return PersonRecord(record_id=record_id, **data)


def _header_text(value, index: int) -> str:
    if value is None or str(value).strip() == "":
        return f"Column {index}"
    return str(value).strip()


def _normalize_header(value: str) -> str:
    text = value.replace("\n", " ").replace("-", " ").replace("/", " ")
    text = text.replace("(", " ").replace(")", " ")
    return " ".join(text.casefold().split())


def _blank(value) -> bool:
    return value is None or str(value).strip() == ""
