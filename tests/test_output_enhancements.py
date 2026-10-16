"""Source provenance, error separation, search budget, checkpoints and workbook output."""

from openpyxl import Workbook, load_workbook

from src.models import PersonRecord
from src.query_planner import FAMILY_HOLDER_ADDRESS, FAMILY_HOLDERS, build_query_plan
from src.scraper.checkpoint import SqliteCheckpoint
from src.scraper.document_investigator import SourceDocument
from src.scraper.excel_io import (
    CORE_OUTPUT_COLUMNS,
    EXTENDED_OUTPUT_COLUMNS,
    load_input,
    run_workbook,
)
from src.scraper.researcher import (
    ResearchStatus,
    ShareholderResearcher,
    _outcome_from_dict,
    _outcome_to_dict,
    research_status_summary,
)
from src.scraper.search_client import SearchPolicy, StaticSearchClient
from src.scraper.source_fetcher import FetchResult
from src.tools.search import SearchResult
from tests.test_obituary_leads import FIXTURE
from tests.test_source_integration import MapSourceFetcher


A1 = "12044 FALCON CREST WAY NORTHRIDGE"
A2 = "CALIFORNIA 91326 U.S.A"


def _record(**overrides):
    data = dict(
        record_id="R",
        name="MITHOO MURLI RAMCHANDANI",
        address1=A1,
        address2=A2,
        country="United States",
    )
    data.update(overrides)
    return PersonRecord(**data)


def _researcher(client, documents, **kwargs):
    return ShareholderResearcher(
        client,
        SearchPolicy(delay_seconds=0, **kwargs),
        source_fetcher=MapSourceFetcher(documents),
        use_result_text_when_unfetched=False,
    )


def test_multiple_sources_and_provenance_are_retained():
    record = _record()
    plans = {plan.strategy: plan.query for plan in build_query_plan(record)}
    search_a = "https://records.example/a?ref=google"
    final_a = "https://mirror.example/a-final"
    url_b = "https://records.example/b"
    url_review = "https://records.example/review"
    client = StaticSearchClient()
    client.register(plans["Q1_FULL_NAME_FULL_ADDRESS"], [SearchResult("A", search_a, "x")])
    client.register(
        plans["Q2_NAME_ADDRESS1"],
        [SearchResult("B", url_b, "x"), SearchResult("Rv", url_review, "x")],
    )
    text = f"MITHOO MURLI RAMCHANDANI\n{A1}\n"
    documents = {
        search_a: SourceDocument(source_url=final_a, title="A", document_type="pdf", text=text),
        url_b: SourceDocument(source_url=url_b, title="B", document_type="html", text=text),
        url_review: SourceDocument(
            source_url=url_review,
            title="Review",
            text="MITHOO MURLI RAMCHANDANI\n12044 FALCON CRST NORTHRIDGE\n",
        ),
    }
    outcome = _researcher(client, documents).research_person(record)
    assert outcome.status == ResearchStatus.STRONG_MATCH
    assert outcome.source_links == [final_a, url_b]
    first = outcome.strong_findings[0]
    assert first.url == final_a
    assert first.original_url == search_a
    assert first.document_type == "pdf"
    assert first.title == "A"
    assert first.group_id and first.evidence_span
    # Review evidence from a separate source is kept, but not as a match.
    assert [item.url for item in outcome.review_findings] == [url_review]


def test_failed_pdf_extraction_is_not_reported_as_no_document_match():
    record = _record()
    url = "https://records.example/scan.pdf"
    client = StaticSearchClient()
    client.register(build_query_plan(record)[0].query, [SearchResult("Scan", url, "x")])

    class FailingPdf:
        def fetch(self, url, title=""):
            return FetchResult(
                ok=False,
                error_code="PDF_EXTRACTION_ERROR",
                error_message="PDF contained no extractable text.",
            )

    outcome = ShareholderResearcher(
        client,
        SearchPolicy(delay_seconds=0),
        source_fetcher=FailingPdf(),
        use_result_text_when_unfetched=False,
    ).research_person(record)
    notes = [note for entry in outcome.log for note in entry.investigation_notes]
    assert any(note.startswith("PDF_EXTRACTION_ERROR") for note in notes)
    assert not any("NO_DOCUMENT_MATCH" in note for note in notes)
    summary = research_status_summary(outcome)
    assert "PDF_EXTRACTION_ERROR x1" in summary
    assert "NO_DOCUMENT_MATCH" not in summary


def test_search_budget_is_respected_and_unexecuted_plans_are_reported():
    record = _record(
        address3="NEAR PARK",
        address4="LOS ANGELES COUNTY",
        pincode="91326",
        scrip="Example Ltd",
        joint1="SUNITA M RAMCHANDANI",
        joint2="KAVITA RAJ SHARMA",
    )
    plans = build_query_plan(record)
    assert len(plans) > 12
    client = StaticSearchClient()
    outcome = _researcher(client, {}).research_person(record)
    holder_plans = [plan for plan in plans if plan.family == FAMILY_HOLDERS]
    joint_address = [plan for plan in plans if plan.family == FAMILY_HOLDER_ADDRESS]
    assert len(holder_plans) == 4
    assert {plan.holder for plan in joint_address} == {2, 3}
    assert all(
        sum(1 for plan in joint_address if plan.holder == holder) <= 12 for holder in (2, 3)
    )
    # 12 primary searches plus the separate per-joint-holder and combination allowances.
    assert len(client.calls) == 12 + len(joint_address) + 4
    assert [entry.query for entry in outcome.log] == client.calls
    executed = {entry.strategy for entry in outcome.log}
    assert {plan.strategy for plan in holder_plans + joint_address} <= executed
    skipped = {item.split(" ")[0] for item in outcome.skipped_strategies}
    assert executed.isdisjoint(skipped)
    assert executed | skipped == {plan.strategy for plan in plans}
    assert all("budget of 12 reached" in item for item in outcome.skipped_strategies)
    assert [item["strategy"] for item in outcome.unexecuted_plans] == [
        item.split(" ")[0] for item in outcome.skipped_strategies
    ]


def test_review_findings_and_leads_survive_checkpoint_roundtrip(tmp_path):
    record = _record(name="PRATAP MALPANI", address1="12 MG ROAD", address2="BANGALORE 560038", country="India")
    url = "https://condolences.example/1"
    client = StaticSearchClient()
    client.register("PRATAP MALPANI obituary", [SearchResult("C", url, "x")])
    store = SqliteCheckpoint(tmp_path / "cp.sqlite")
    documents = {
        url: SourceDocument(
            source_url=url, title="C", text=FIXTURE.read_text(encoding="utf-8")
        )
    }
    first = ShareholderResearcher(
        client,
        SearchPolicy(delay_seconds=0),
        checkpoint=store,
        source_fetcher=MapSourceFetcher(documents),
        use_result_text_when_unfetched=False,
    ).research_person(record)
    assert first.obituary_leads

    class Exploding:
        def search(self, query):
            raise AssertionError("finished rows must not search again")

    resumed = ShareholderResearcher(
        Exploding(), SearchPolicy(delay_seconds=0), checkpoint=store
    ).research_person(record)
    assert resumed.obituary_leads[0].to_dict() == first.obituary_leads[0].to_dict()
    assert resumed.skipped_strategies == first.skipped_strategies

    restored = _outcome_from_dict(_outcome_to_dict(first))
    assert restored.obituary_leads[0].companies == ["VASP TRADING (S) PTE LTD"]


def _workbook(path, rows):
    book = Workbook()
    sheet = book.active
    sheet.append(
        [
            "CL-ID/Folio No.",
            "Name of the Holder",
            "Second Holder (Joint 1)",
            "Third Holder (Joint-2)",
            "Address-1",
            "Address-2",
            "Address-3",
            "Address-4",
            "Pincode",
            " Country",
        ]
    )
    for row in rows:
        sheet.append(row)
    book.save(path)
    book.close()


def test_workbook_exports_manual_review_leads_and_keeps_existing_columns(tmp_path):
    source = tmp_path / "in.xlsx"
    destination = tmp_path / "out.xlsx"
    _workbook(
        source,
        [
            ["F1", "MITHOO MURLI RAMCHANDANI", "SUNITA M RAMCHANDANI", None, A1, A2, None, None, "0", "United States"],
            ["F2", "PRATAP MALPANI", None, None, "12 MG ROAD INDIRANAGAR", "BANGALORE 560038", None, None, "0", "India"],
        ],
    )
    rows = load_input(source)[2]
    review_url = "https://records.example/review"
    obit_url = "https://condolences.example/1"
    client = StaticSearchClient()
    client.register(build_query_plan(rows[0].record)[0].query, [SearchResult("R", review_url, "x")])
    client.register("PRATAP MALPANI obituary", [SearchResult("C", obit_url, "x")])
    documents = {
        review_url: SourceDocument(
            source_url=review_url,
            title="Dividend list",
            document_type="pdf",
            text="MITHOO MURLI RAMCHANDANI\tSUNITA M RAMCHANDANI\t12044 FALCON CRST NORTHRIDGE\n",
        ),
        obit_url: SourceDocument(
            source_url=obit_url,
            title="Condolences",
            document_type="html",
            text=FIXTURE.read_text(encoding="utf-8"),
        ),
    }
    run_workbook(
        source,
        destination,
        client,
        SearchPolicy(delay_seconds=0),
        source_fetcher=MapSourceFetcher(documents),
    )

    saved = load_workbook(destination)
    main = saved["Sheet"]
    headers = [cell.value for cell in main[1]]
    assert headers[10:21] == CORE_OUTPUT_COLUMNS
    assert headers[21:] == EXTENDED_OUTPUT_COLUMNS
    review_row = {h: c.value for h, c in zip(headers, main[2])}
    obit_row = {h: c.value for h, c in zip(headers, main[3])}

    assert review_row["Match Status"] == "MANUAL_REVIEW"
    assert review_row["Matched Name"] in (None, "")
    assert review_row["Source Links"] in (None, "")
    assert review_row["Manual Review Sources"] == review_url
    assert "EXACT_NAME + ADDRESS_90_TO_95" in review_row["Manual Review Reason"]
    assert review_row["Match Basis"] == "EXACT_NAME + ADDRESS_90_TO_95"
    assert review_row["Name Match Classification"] == "EXACT"
    assert 90 <= review_row["Address 1 Match %"] < 95
    assert "joint1 'SUNITA M RAMCHANDANI': EXACT in same record" in review_row["Joint Holder Evidence"]
    assert "[MANUAL_REVIEW] " + review_url in review_row["Source Evidence"]
    assert "91326" in review_row["Derived Postal Fragments"]
    assert "MANUAL_REVIEW_CANDIDATE x1" in review_row["Research Status / Errors"]

    assert obit_row["Match Status"] == "NO_STRONG_MATCH"
    assert obit_url in obit_row["Obituary / Condolence Lead"]
    assert "PARENTS" in obit_row["Family/Relationship Evidence"]
    assert "BANGALORE, INDIA" in obit_row["Location Evidence"]
    assert "VASP TRADING (S) PTE LTD" in obit_row["Company Evidence"]
    assert "relationship not established" in obit_row["Company Evidence"]

    evidence = list(saved["Source Evidence"].iter_rows(values_only=True))
    assert evidence[0][2] == "Classification"
    assert [(row[0], row[2], row[14]) for row in evidence[1:]] == [("2", "MANUAL_REVIEW", review_url)]

    leads = list(saved["Obituary Leads"].iter_rows(values_only=True))
    assert len(leads) == 2 and leads[1][2] == obit_url
    assert "SUDARSHAN" in leads[1][11]

    log_rows = list(saved["Research Log"].iter_rows(values_only=True))
    assert log_rows[0][-2:] == ("Investigation Notes", "Google Results")
    executed = {row[3] for row in log_rows[1:] if row[6] != "NOT_EXECUTED"}
    assert "PRATAP MALPANI obituary" in executed
    # Planned-but-skipped searches are logged with a reason, never as completed.
    for row in log_rows[1:]:
        if row[6] == "NOT_EXECUTED":
            assert row[4] is None and row[5] is None
            assert row[7]
    saved.close()
