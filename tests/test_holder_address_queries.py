"""Per-holder name + address searches, their budgets, logging and resume."""

from pathlib import Path

import pytest
from openpyxl import load_workbook

from src.models import PersonRecord
from src.normalize import normalize_text
from src.query_planner import (
    FAMILY_CORE,
    FAMILY_HOLDER_ADDRESS,
    FAMILY_HOLDERS,
    FAMILY_OBITUARY,
    build_query_plan,
)
from src.scraper.checkpoint import SqliteCheckpoint
from src.scraper.excel_io import load_input, run_workbook
from src.scraper.researcher import ResearchStatus, ShareholderResearcher
from src.scraper.search_client import SearchPolicy, SearchStatus, StaticSearchClient
from tests.test_output_enhancements import _workbook
from tests.test_source_integration import MapSourceFetcher


A, B, C = "MITHOO MURLI RAMCHANDANI", "SUNITA M RAMCHANDANI", "KAVITA RAJ SHARMA"
A1 = "12044 FALCON CREST WAY NORTHRIDGE"
A2 = "CALIFORNIA 91326 U.S.A"
A3 = "NEAR PARK"
A4 = "LOS ANGELES COUNTY"
PIN = "91326"
COUNTRY = "United States"

ROW27_A1 = "329 RIVER VALLEY ROAD 1802 SINGAPORE"


def _record(**overrides):
    data = dict(
        record_id="T",
        name=A,
        joint1=B,
        joint2=C,
        address1=A1,
        address2=A2,
        address3=A3,
        address4=A4,
        pincode=PIN,
        country=COUNTRY,
    )
    data.update(overrides)
    return PersonRecord(**data)


def _row27():
    return PersonRecord(
        record_id="27",
        name="VIBHA MALPANI",
        joint1="PRATAP MALPANI",
        address1=ROW27_A1,
        address2="923",
        country="Singapore",
    )


def _address_plans(record, holder):
    """strategy (holder suffix removed) -> query, for one holder's address variants."""
    out = {}
    for plan in build_query_plan(record):
        if plan.strategy.startswith("Q_POSTAL_"):
            continue
        if holder == 1 and plan.family == FAMILY_CORE and plan.strategy != "Q5_NAME_SCRIP":
            out[plan.strategy] = plan.query
        elif plan.family == FAMILY_HOLDER_ADDRESS and plan.holder == holder:
            out[plan.strategy.removesuffix(f"_HOLDER{holder}")] = plan.query
    return out


def _expected_variants(holder_name):
    return {
        "Q1_FULL_NAME_FULL_ADDRESS": f"{holder_name} {A1} {A2} {A3} {A4} {PIN}",
        "Q_NAME_ADDRESS1_ADDRESS2": f"{holder_name} {A1} {A2}",
        "Q2_NAME_ADDRESS1": f"{holder_name} {A1}",
        "Q3_NAME_ADDRESS2": f"{holder_name} {A2}",
        "Q_NAME_ADDRESS3": f"{holder_name} {A3}",
        "Q_NAME_ADDRESS4": f"{holder_name} {A4}",
        "Q_NAME_ADDRESS_FRAGMENT": f"{holder_name} FALCON CREST WAY NORTHRIDGE",
        "Q_NAME_PINCODE": f"{holder_name} {PIN}",
        "Q4_NAME_COUNTRY": f"{holder_name} {COUNTRY}",
    }


def _researcher(client, **kwargs):
    checkpoint = kwargs.pop("checkpoint", None)
    return ShareholderResearcher(
        client,
        SearchPolicy(delay_seconds=0, **kwargs),
        checkpoint=checkpoint,
        source_fetcher=MapSourceFetcher({}),
        use_result_text_when_unfetched=False,
    )


# ---------------------------------------------------------------- planning


def test_three_holders_each_get_every_address_variant_independently():
    record = _record()
    for holder, holder_name in ((1, A), (2, B), (3, C)):
        assert _address_plans(record, holder) == _expected_variants(holder_name)
    holder_queries = [
        plan.query for plan in build_query_plan(record) if plan.family == FAMILY_HOLDER_ADDRESS
    ]
    # The primary name is never copied into a joint holder's address query.
    assert all(A not in query for query in holder_queries)


def test_derived_postal_code_variant_is_applied_to_each_holder():
    plans = {plan.strategy: plan.query for plan in build_query_plan(_record())}
    assert plans["Q_POSTAL_NAME_LOCATION"] == f"{A} CALIFORNIA {PIN}"
    assert plans["Q_POSTAL_NAME_LOCATION_HOLDER2"] == f"{B} CALIFORNIA {PIN}"
    assert plans["Q_POSTAL_NAME_LOCATION_HOLDER3"] == f"{C} CALIFORNIA {PIN}"


def test_primary_address_strategies_are_unchanged_by_joint_holders():
    with_joints = [p for p in build_query_plan(_record()) if p.family == FAMILY_CORE]
    alone = [p for p in build_query_plan(_record(joint1="", joint2="")) if p.family == FAMILY_CORE]
    assert with_joints == alone
    assert [p.strategy for p in with_joints][:3] == [
        "Q1_FULL_NAME_FULL_ADDRESS",
        "Q_NAME_ADDRESS1_ADDRESS2",
        "Q2_NAME_ADDRESS1",
    ]


def test_joint_holder_address_queries_follow_the_primary_core_block():
    families = [plan.family for plan in build_query_plan(_record())]
    first = families.index(FAMILY_HOLDER_ADDRESS)
    assert set(families[:first]) == {FAMILY_CORE}
    last = len(families) - 1 - families[::-1].index(FAMILY_HOLDER_ADDRESS)
    assert set(families[first : last + 1]) == {FAMILY_HOLDER_ADDRESS}
    # Pair, all-holder and obituary searches are all still planned.
    assert {FAMILY_HOLDERS, FAMILY_OBITUARY} <= set(families)


def test_row27_searches_the_address_for_both_holders():
    plans = {plan.strategy: plan.query for plan in build_query_plan(_row27())}
    for suffix, holder_name in (("", "VIBHA MALPANI"), ("_HOLDER2", "PRATAP MALPANI")):
        assert plans[f"Q1_FULL_NAME_FULL_ADDRESS{suffix}"] == f"{holder_name} {ROW27_A1} 923"
        assert plans[f"Q2_NAME_ADDRESS1{suffix}"] == f"{holder_name} {ROW27_A1}"
        assert plans[f"Q3_NAME_ADDRESS2{suffix}"] == f"{holder_name} 923"
        assert plans[f"Q_NAME_ADDRESS_FRAGMENT{suffix}"] == (
            f"{holder_name} RIVER VALLEY ROAD 1802 SINGAPORE"
        )
        assert plans[f"Q4_NAME_COUNTRY{suffix}"] == f"{holder_name} Singapore"
    # Existing joint, pair and obituary searches remain.
    assert plans["Q_JOINT_PRIMARY_JOINT_ADDRESS"].startswith("VIBHA MALPANI PRATAP MALPANI")
    assert plans["Q_HOLDER_PAIR_1_2"] == "VIBHA MALPANI PRATAP MALPANI"
    assert plans["Q_OBIT_NAME_OBITUARY"] == "VIBHA MALPANI obituary"


def test_row27_from_the_real_input_workbook():
    source = Path(__file__).resolve().parents[1] / "input.xlsx"
    if not source.exists():
        pytest.skip("input.xlsx not present")
    rows = load_input(source)[2]
    record = next(row.record for row in rows if row.record.name.strip() == "VIBHA MALPANI")
    assert record.joint1.strip() == "PRATAP MALPANI"
    queries = {plan.query for plan in build_query_plan(record)}
    for holder_name in ("VIBHA MALPANI", "PRATAP MALPANI"):
        assert f"{holder_name} {ROW27_A1} 923" in queries
        assert f"{holder_name} 923" in queries


def test_one_holder_has_no_joint_address_queries():
    plans = build_query_plan(_record(joint1="", joint2=""))
    assert not any(plan.family == FAMILY_HOLDER_ADDRESS for plan in plans)
    assert _address_plans(_record(joint1="", joint2=""), 1) == _expected_variants(A)


def test_two_holders_and_blank_second_or_third_holder():
    two = _record(joint2="")
    assert _address_plans(two, 2) == _expected_variants(B)
    assert _address_plans(two, 3) == {}

    third_only = _record(joint1="   ", joint2=C)
    assert _address_plans(third_only, 2) == {}
    assert _address_plans(third_only, 3) == _expected_variants(C)
    for plan in build_query_plan(third_only):
        assert plan.query == plan.query.strip() and "  " not in plan.query


def test_empty_address_fields_are_skipped_for_every_holder():
    record = _record(address2="", address3="", address4="", pincode="", country="")
    for holder, holder_name in ((1, A), (2, B), (3, C)):
        assert _address_plans(record, holder) == {
            "Q1_FULL_NAME_FULL_ADDRESS": f"{holder_name} {A1}",
            "Q_NAME_ADDRESS_FRAGMENT": f"{holder_name} FALCON CREST WAY NORTHRIDGE",
        }


def test_missing_primary_still_searches_joint_holders_with_the_address():
    record = _record(name="")
    plans = build_query_plan(record)
    assert _address_plans(record, 2) == _expected_variants(B)
    assert _address_plans(record, 3) == _expected_variants(C)
    # No address-only or company-only search masquerades as a holder search.
    assert not any(plan.family == FAMILY_CORE for plan in plans)


def test_deduplication_keeps_holder_specific_queries():
    plans = build_query_plan(_record())
    keys = [normalize_text(plan.query) for plan in plans]
    assert len(keys) == len(set(keys))
    # Holders sharing a surname are distinct people with distinct queries.
    assert _address_plans(_record(), 2)["Q2_NAME_ADDRESS1"] != _address_plans(_record(), 1)["Q2_NAME_ADDRESS1"]

    # A joint field repeating the primary (any case/spacing) adds nothing.
    duplicate = _record(joint1=f"  {A.lower()} ", joint2="")
    assert not any(plan.family == FAMILY_HOLDER_ADDRESS for plan in build_query_plan(duplicate))
    same_joint = _record(joint2=B.title())
    assert _address_plans(same_joint, 3) == {}
    assert _address_plans(same_joint, 2) == _expected_variants(B)


# --------------------------------------------------------------- execution


def test_primary_budget_cannot_starve_joint_holder_address_searches():
    record = _record()
    client = StaticSearchClient()
    outcome = _researcher(client, max_searches_per_person=1).research_person(record)
    for holder in (2, 3):
        for query in _address_plans(record, holder).values():
            assert query in client.calls
    executed = {entry.strategy for entry in outcome.log}
    assert "Q1_FULL_NAME_FULL_ADDRESS" in executed
    assert "Q2_NAME_ADDRESS1" not in executed
    reasons = {item["strategy"]: item["reason"] for item in outcome.unexecuted_plans}
    assert reasons["Q2_NAME_ADDRESS1"] == "per-person search budget of 1 reached"


def test_per_holder_limit_reports_every_unexecuted_search():
    record = _record()
    client = StaticSearchClient()
    outcome = _researcher(client, max_searches_per_additional_holder=2).research_person(record)
    unexecuted = {item["strategy"]: item for item in outcome.unexecuted_plans}
    for holder in (2, 3):
        plans = [
            plan
            for plan in build_query_plan(record)
            if plan.family == FAMILY_HOLDER_ADDRESS and plan.holder == holder
        ]
        ran = [plan for plan in plans if plan.query in client.calls]
        assert [plan.strategy for plan in ran] == [plan.strategy for plan in plans[:2]]
        for plan in plans[2:]:
            item = unexecuted[plan.strategy]
            assert item["query"] == plan.query
            assert item["reason"] == f"holder {holder} address search budget of 2 reached"
    executed = {entry.strategy for entry in outcome.log}
    assert executed.isdisjoint(unexecuted)
    assert executed | set(unexecuted) == {plan.strategy for plan in build_query_plan(record)}


def test_blocked_search_marks_remaining_plans_not_executed():
    record = _record(joint2="")
    plans = build_query_plan(record)
    client = StaticSearchClient()
    client.register(plans[0].query, status=SearchStatus.BLOCKED, message="captcha shown")
    outcome = _researcher(client, max_consecutive_failures=1).research_person(record)
    assert [entry.strategy for entry in outcome.log] == [plans[0].strategy]
    assert outcome.log[0].status == SearchStatus.BLOCKED
    assert [item["strategy"] for item in outcome.unexecuted_plans] == [
        plan.strategy for plan in plans[1:]
    ]
    assert all(item["reason"] == "search halted" for item in outcome.unexecuted_plans)
    assert outcome.status == ResearchStatus.BLOCKED


def test_research_log_lists_executed_and_unexecuted_queries(tmp_path):
    source, destination = tmp_path / "in.xlsx", tmp_path / "out.xlsx"
    _workbook(source, [["F27", "VIBHA MALPANI", "PRATAP MALPANI", None, ROW27_A1, "923", None, None, "0", "Singapore"]])
    record = load_input(source)[2][0].record
    plans = build_query_plan(record)
    run_workbook(
        source,
        destination,
        StaticSearchClient(),
        SearchPolicy(delay_seconds=0, max_searches_per_person=3),
        source_fetcher=MapSourceFetcher({}),
    )
    book = load_workbook(destination, read_only=True)
    rows = list(book["Research Log"].iter_rows(values_only=True))[1:]
    book.close()
    logged = {(row[2], row[3]): row for row in rows}
    assert len(logged) == len(rows) == len(plans)
    for plan in plans:
        row = logged[(plan.strategy, plan.query)]
        if row[6] == "NOT_EXECUTED":
            assert row[4] is None and "budget of 3 reached" in row[7]
        else:
            assert row[6] == "OK" and row[4] == 0
    executed = [row[3] for row in rows if row[6] == "OK"]
    assert "PRATAP MALPANI 329 RIVER VALLEY ROAD 1802 SINGAPORE 923" in executed
    assert "PRATAP MALPANI 923" in executed


class _InterruptAfter(StaticSearchClient):
    def __init__(self, limit):
        super().__init__()
        self.limit = limit

    def search(self, query):
        if len(self.calls) >= self.limit:
            raise KeyboardInterrupt
        return super().search(query)


def test_interrupted_run_writes_unexecuted_searches_and_resumes(tmp_path):
    source, destination = tmp_path / "in.xlsx", tmp_path / "out.xlsx"
    _workbook(
        source,
        [
            ["F27", "VIBHA MALPANI", "PRATAP MALPANI", None, ROW27_A1, "923", None, None, "0", "Singapore"],
            ["F28", "ABDUL SAMAD KHAN", None, None, "PO BOX 967 RIYADH 11421", "SAUDI ARABIA", None, None, "0", "Saudi Arabia"],
        ],
    )
    first, second = (row.record for row in load_input(source)[2])
    store = SqliteCheckpoint(tmp_path / "cp.sqlite")
    policy = SearchPolicy(delay_seconds=0)
    with pytest.raises(KeyboardInterrupt):
        run_workbook(source, destination, _InterruptAfter(3), policy, checkpoint=store,
                     source_fetcher=MapSourceFetcher({}))

    book = load_workbook(destination, read_only=True)
    rows = list(book["Research Log"].iter_rows(values_only=True))[1:]
    book.close()
    first_rows = [row for row in rows if row[0] == first.record_id and row[2]]
    assert [row[6] for row in first_rows[:3]] == ["OK", "OK", "OK"]
    assert all(row[6] == "NOT_EXECUTED" for row in first_rows[3:])
    assert all("run interrupted" in row[7] for row in first_rows[3:])
    assert len(first_rows) == len(build_query_plan(first))
    second_rows = [row for row in rows if row[0] == second.record_id and row[2]]
    assert len(second_rows) == len(build_query_plan(second))
    assert all(row[6] == "NOT_EXECUTED" for row in second_rows)
    status_row = next(row for row in rows if row[0] == second.record_id and not row[2])
    assert status_row[6] == "ERROR" and status_row[7].startswith("INTERRUPTED")

    # Resume: nothing that already ran is searched again.
    resumed = StaticSearchClient()
    outcomes = run_workbook(source, destination, resumed, policy, checkpoint=store,
                            source_fetcher=MapSourceFetcher({}))
    first_plans = [plan.query for plan in build_query_plan(first)]
    assert resumed.calls[: len(first_plans) - 3] == first_plans[3:]
    assert not set(first_plans[:3]) & set(resumed.calls)
    assert len(outcomes[0].log) == len(first_plans)
    assert outcomes[0].unexecuted_plans == []


def test_interrupted_outcome_is_not_reported_as_complete():
    record = _row27()
    researcher = _researcher(_InterruptAfter(2))
    with pytest.raises(KeyboardInterrupt):
        researcher.research_person(record)
    outcome = researcher.interrupted_outcome(record)
    assert outcome.status == ResearchStatus.ERROR
    assert outcome.message.startswith("INTERRUPTED")
    assert len(outcome.log) == 2
    assert len(outcome.unexecuted_plans) == len(build_query_plan(record)) - 2
