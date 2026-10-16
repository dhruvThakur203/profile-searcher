from openpyxl import Workbook, load_workbook

from src.models import PersonRecord
from src.query_planner import build_query_plan
from src.scraper.excel_io import (
    EXTENDED_OUTPUT_COLUMNS,
    MOBILE_SEPARATOR,
    SOURCE_LINK_SEPARATOR,
    load_input,
    run_workbook,
)
from src.scraper.search_client import SearchPolicy, StaticSearchClient
from src.tools.search import SearchResult


def test_excel_roundtrip_preserves_rows_and_input_file(tmp_path):
    source = tmp_path / "shareholders.xlsx"
    destination = tmp_path / "shareholders_out.xlsx"
    built = Workbook()
    sheet = built.active
    sheet.title = "Holders"
    sheet.append(["Name", "Address 1", "Address 2", "Country", "Scrip / Company", "Notes"])
    sheet.append(
        [
            "Dhruv Singh Thakur",
            "123 ABC Road",
            "Singapore 123456",
            "Singapore",
            "Example Ltd",
            "keep-me",
        ]
    )
    sheet.append(
        [
            "Meera Lal",
            "123 ABC Road",
            "Singapore 123456",
            "Singapore",
            "Example Ltd",
            "also-keep",
        ]
    )
    built.save(source)
    built.close()
    original = source.read_bytes()

    first = PersonRecord(
        record_id="2",
        name="Dhruv Singh Thakur",
        address1="123 ABC Road",
        address2="Singapore 123456",
        country="Singapore",
        scrip="Example Ltd",
    )
    second = PersonRecord(
        record_id="3",
        name="Meera Lal",
        address1="123 ABC Road",
        address2="Singapore 123456",
        country="Singapore",
        scrip="Example Ltd",
    )
    client = StaticSearchClient()
    client.register(
        build_query_plan(first)[0].query,
        [
            SearchResult(
                title="Source A",
                url="https://source1.example/a",
                snippet=(
                    "Dhruv Singh Thakur 123 ABC Road Singapore 123456 "
                    "mobile 9999999999 or 8888888888"
                ),
            ),
            SearchResult(
                title="Source B",
                url="https://source2.example/b",
                snippet=(
                    "Dhruv Singh Thakur 123 ABC Road Singapore 123456 "
                    "email person@example.com mobile 9999999999"
                ),
            ),
        ],
    )
    client.register(
        build_query_plan(second)[0].query,
        [
            SearchResult(
                title="Source C",
                url="https://source3.example/c",
                snippet="Meera Lal 123 ABC Road Singapore 123456 email meera@example.com",
            )
        ],
    )

    run_workbook(source, destination, client, SearchPolicy(delay_seconds=0))

    assert source.read_bytes() == original
    saved = load_workbook(destination)
    holders = saved["Holders"]
    headers = [cell.value for cell in holders[1]]
    assert headers[:6] == [
        "Name",
        "Address 1",
        "Address 2",
        "Country",
        "Scrip / Company",
        "Notes",
    ]
    # Existing output columns keep their names and positions; new columns are
    # only ever appended after them.
    assert headers[6:17] == [
        "Matched Name",
        "Mobile",
        "Email",
        "Source Links",
        "Evidence",
        "Match Status",
        "Address 1 Match %",
        "Address 2 Match %",
        "Address 3 Match %",
        "Address 4 Match %",
        "Search Strategies Used",
    ]
    assert headers[17:] == EXTENDED_OUTPUT_COLUMNS
    assert holders.max_row == 3
    first_row = [cell.value for cell in holders[2]]
    second_row = [cell.value for cell in holders[3]]
    assert first_row[0] == "Dhruv Singh Thakur"
    assert first_row[5] == "keep-me"
    assert first_row[6] == "Dhruv Singh Thakur"
    assert first_row[7] == MOBILE_SEPARATOR.join(["9999999999", "8888888888"])
    assert first_row[8] == "person@example.com"
    assert first_row[9] == SOURCE_LINK_SEPARATOR.join(
        ["https://source1.example/a", "https://source2.example/b"]
    )
    assert "meera@example.com" not in str(first_row[8])
    assert first_row[11] == "STRONG_MATCH"
    assert second_row[0] == "Meera Lal"
    assert second_row[5] == "also-keep"
    assert second_row[7] in (None, "")
    assert second_row[8] == "meera@example.com"
    assert second_row[9] == "https://source3.example/c"
    assert "9999999999" not in "".join(str(value) for value in second_row)
    assert second_row[11] == "STRONG_MATCH"

    log = saved["Research Log"]
    queries = [row[3] for row in log.iter_rows(min_row=2, values_only=True)]
    assert build_query_plan(first)[0].query in queries
    assert build_query_plan(second)[0].query in queries
    saved.close()


def _shareholder_row(path, pincode):
    book = Workbook()
    sheet = book.active
    sheet.append(
        [
            "Name of the Holder",
            "Address-1",
            "Address-2",
            "Address-3",
            "Address-4",
            "Pincode",
            " Country",
        ]
    )
    sheet.append(
        [
            "CIDAMBI S RAVI",
            "16708 AMBERWOOD WAY",
            "CERRITOS CA 90703 USA",
            None,
            None,
            pincode,
            "United States",
        ]
    )
    book.save(path)
    book.close()
    return load_input(path)[2][0].record


def test_placeholder_pincode_is_treated_as_empty(tmp_path):
    record = _shareholder_row(tmp_path / "holders.xlsx", "0")
    assert record.pincode == ""
    assert record.name == "CIDAMBI S RAVI"
    assert record.address1 == "16708 AMBERWOOD WAY"
    assert record.address2 == "CERRITOS CA 90703 USA"
    assert record.address3 == ""
    assert record.address4 == ""
    assert record.country == "United States"


def test_genuine_pincode_is_retained(tmp_path):
    record = _shareholder_row(tmp_path / "holders.xlsx", "999999")
    assert record.pincode == "999999"


def test_q1_omits_placeholder_pincode(tmp_path):
    record = _shareholder_row(tmp_path / "holders.xlsx", "0")
    full = next(
        plan.query
        for plan in build_query_plan(record)
        if plan.strategy == "Q1_FULL_NAME_FULL_ADDRESS"
    )
    assert full == "CIDAMBI S RAVI 16708 AMBERWOOD WAY CERRITOS CA 90703 USA"
    assert not full.endswith(" 0")


def test_name_pincode_query_is_not_generated_for_placeholder(tmp_path):
    record = _shareholder_row(tmp_path / "holders.xlsx", "0")
    strategies = [plan.strategy for plan in build_query_plan(record)]
    assert "Q_NAME_PINCODE" not in strategies
