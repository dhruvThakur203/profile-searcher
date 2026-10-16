"""Google AI Overview capture, leads, Excel export, and identity safety."""

from pathlib import Path

from openpyxl import Workbook, load_workbook

from src.models import PersonRecord
from src.query_planner import build_query_plan
from src.scraper.ai_overview import (
    STATUS_BLOCKED,
    STATUS_CAPTURE_FAILED,
    STATUS_NOT_PRESENT,
    STATUS_PRESENT,
    STATUS_UNAVAILABLE,
    AiOverviewCapture,
    AiOverviewLogEntry,
    AiOverviewSource,
    AiContactLead,
    capture_from_playwright_payload,
    dedupe_contact_leads,
    extract_contact_leads_from_overview,
    parse_ai_overview_from_html,
    revalidate_log_entry,
    sanitize_stored_ai_evidence,
)
from src.scraper.checkpoint import SqliteCheckpoint
from src.scraper.excel_io import (
    AI_CONTACT_LEAD_COLUMNS,
    AI_OVERVIEW_LOG_COLUMNS,
    EXTENDED_OUTPUT_COLUMNS,
    run_workbook,
)
from src.scraper.researcher import (
    ResearchStatus,
    ShareholderResearcher,
    _outcome_from_dict,
    _outcome_to_dict,
)
from src.scraper.search_client import SearchPolicy, SearchStatus, StaticSearchClient
from src.tools.search import SearchResult

FIXTURES = Path(__file__).resolve().parent / "fixtures" / "google"
OVERVIEW_HTML = (FIXTURES / "ai_overview_visweswaran.html").read_text(encoding="utf-8")
SNIPPET_HTML = (FIXTURES / "serp_snippet_only.html").read_text(encoding="utf-8")
TRUNCATED_HTML = (FIXTURES / "ai_overview_truncated.html").read_text(encoding="utf-8")
UNAVAILABLE_HTML = (FIXTURES / "ai_overview_google_unavailable.html").read_text(encoding="utf-8")
ORGANIC_PRIVACY_HTML = (FIXTURES / "serp_organic_with_privacy_no_overview.html").read_text(
    encoding="utf-8"
)
NO_SOURCES_HTML = (FIXTURES / "ai_overview_present_no_sources.html").read_text(encoding="utf-8")
JSNAME_HTML = (FIXTURES / "ai_overview_live_jsname.html").read_text(encoding="utf-8")
PROPERTY_JSNAME_HTML = (FIXTURES / "ai_overview_property_jsname.html").read_text(encoding="utf-8")
LIVE_FALSE_POSITIVE_TEXT = (
    "An AI Overview is not available for this searchCan't generate an AI overview right now. "
    "Try again later.AI Overview .o3lpyf{background-color:var(--BRLwE);}"
)


def _record(**overrides):
    data = dict(
        record_id="9",
        name="VISWESWARAN NAVARATNAM",
        address1="12 RESEARCH PARK",
        address2="PENANG",
        country="Malaysia",
    )
    data.update(overrides)
    return PersonRecord(**data)


def test_live_workbook_unavailable_message_is_not_present():
    capture = parse_ai_overview_from_html(UNAVAILABLE_HTML)
    assert capture.status == STATUS_UNAVAILABLE
    assert capture.raw_text == ""
    assert not capture.sources


def test_organic_serp_with_privacy_is_not_present():
    capture = parse_ai_overview_from_html(ORGANIC_PRIVACY_HTML)
    assert capture.status == STATUS_NOT_PRESENT


def test_css_with_ai_overview_label_in_snippet_is_not_present():
    html = (
        "<html><body><div class='g MjjYud'>"
        "<span>AI Overview</span><style>.o3lpyf{background-color:var(--BRLwE);}</style>"
        "<a href='https://www.zillow.com/x'>Zillow</a></div></body></html>"
    )
    capture = parse_ai_overview_from_html(html)
    assert capture.status == STATUS_NOT_PRESENT


def test_stale_unavailable_copy_with_real_overview_is_present():
    html = """
    <html><body><div id="cnt"><div jsname="ZLxsqf">
      <span>An AI Overview is not available for this search</span>
      <span>Can't generate an AI overview right now. Try again later.</span>
      <div><span>AI Overview</span><p>Jane Example email help@example.org for media inquiries.</p></div>
      <a href="https://example.org/profile">Example profile</a>
    </div></div></body></html>
    """
    capture = parse_ai_overview_from_html(html)
    assert capture.status == STATUS_PRESENT
    assert "help@example.org" in capture.raw_text


def test_current_jsname_fixture_without_sgcop_is_present():
    capture = parse_ai_overview_from_html(JSNAME_HTML)
    assert capture.status == STATUS_PRESENT
    assert "nava@usm.my" in capture.raw_text
    assert len(capture.sources) >= 1
    assert "policies.google.com" not in " ".join(item.url for item in capture.sources)


def test_legacy_sgcop_fixture_remains_present():
    capture = parse_ai_overview_from_html(OVERVIEW_HTML)
    assert capture.status == STATUS_PRESENT


def test_html_fallback_when_js_payload_not_present():
    capture_js = capture_from_playwright_payload(
        {"state": "not_present", "text": "", "sources": [], "truncated": False, "notes": ""}
    )
    assert capture_js.status == STATUS_NOT_PRESENT
    capture_html = parse_ai_overview_from_html(JSNAME_HTML)
    assert capture_html.status == STATUS_PRESENT
    assert "nava@usm.my" in capture_html.raw_text


def test_property_style_overview_without_contact_leads():
    capture = parse_ai_overview_from_html(PROPERTY_JSNAME_HTML)
    assert capture.status == STATUS_PRESENT
    assert "Cerritos" in capture.raw_text
    leads = extract_contact_leads_from_overview(
        capture,
        excel_row=2,
        folio_number="F",
        holder_associated="CIDAMBI S RAVI",
        search_query="property query",
    )
    assert leads == []


def test_parse_valid_ai_overview_with_sources():
    capture = parse_ai_overview_from_html(OVERVIEW_HTML)
    assert capture.status == STATUS_PRESENT
    assert "nava@usm.my" in capture.raw_text
    assert len(capture.sources) == 2
    assert "policies.google.com" not in " ".join(item.url for item in capture.sources)


def test_genuine_overview_without_source_links():
    capture = parse_ai_overview_from_html(NO_SOURCES_HTML)
    assert capture.status == STATUS_PRESENT
    assert capture.sources == []
    assert "communications office" in capture.raw_text


def test_blocked_page_status():
    html = "<html><body><form id='captcha-form'>unusual traffic</form></body></html>"
    capture = parse_ai_overview_from_html(html)
    assert capture.status == STATUS_BLOCKED


def test_invalid_css_phone_candidates_rejected():
    capture = AiOverviewCapture(
        status=STATUS_PRESENT,
        raw_text=(
            "Contact details may include 2147483648 or 2654435769 in broken markup "
            "var(--token) .class{height:22px}"
        ),
    )
    leads = extract_contact_leads_from_overview(
        capture,
        excel_row=1,
        folio_number="",
        holder_associated="X",
        search_query="q",
    )
    assert leads == []


def test_valid_email_and_international_phone_from_genuine_overview():
    capture = parse_ai_overview_from_html(OVERVIEW_HTML)
    leads = extract_contact_leads_from_overview(
        capture,
        excel_row=1,
        folio_number="",
        holder_associated="VISWESWARAN NAVARATNAM",
        search_query="q",
    )
    assert any(lead.email == "nava@usm.my" for lead in leads)
    phones = [lead.phone for lead in leads if lead.phone]
    assert phones
    assert any("657 9022" in phone or "+60" in phone for phone in phones)
    assert all(lead.institution != "var" for lead in leads)


def test_no_contacts_from_unavailable_capture():
    capture = parse_ai_overview_from_html(UNAVAILABLE_HTML)
    assert extract_contact_leads_from_overview(
        capture,
        excel_row=1,
        folio_number="",
        holder_associated="X",
        search_query="q",
    ) == []


def test_unavailable_log_entry_has_blank_raw_text():
    from src.scraper.ai_overview import build_log_entry

    entry = build_log_entry(
        capture=parse_ai_overview_from_html(UNAVAILABLE_HTML),
        excel_row=2,
        folio_number="F1",
        primary_holder="A",
        holder_associated="A",
        query_strategy="Q1",
        search_query="query",
        capture_timestamp="t",
    )
    assert entry.overview_status == STATUS_UNAVAILABLE
    assert entry.raw_overview_text == ""
    assert "not available" in entry.extraction_notes.casefold()


def test_sanitize_legacy_false_positive_workbook_capture():
    bad = AiOverviewLogEntry(
        excel_row=2,
        folio_number="F",
        primary_holder="CIDAMBI S RAVI",
        holder_associated="CIDAMBI S RAVI",
        query_strategy="Q1",
        search_query="CIDAMBI S RAVI CERRITOS CA 90703 USA",
        capture_timestamp="t",
        overview_status=STATUS_PRESENT,
        raw_overview_text=LIVE_FALSE_POSITIVE_TEXT,
        source_titles=["Zillow", "Privacy Policy"],
        source_urls=["https://policies.google.com/privacy?hl=en-IN"],
    )
    bad_leads = [
        AiContactLead(
            phone="2147483648",
            institution="var",
            holder_associated="CIDAMBI S RAVI",
            search_query="CIDAMBI S RAVI CERRITOS CA 90703 USA",
            supporting_evidence_text=LIVE_FALSE_POSITIVE_TEXT,
        )
    ]
    log, leads = sanitize_stored_ai_evidence([bad], bad_leads)
    assert log[0].overview_status in {STATUS_UNAVAILABLE, STATUS_CAPTURE_FAILED}
    assert log[0].raw_overview_text == ""
    assert leads == []


def test_checkpoint_roundtrip_revalidates_false_positive():
    record = _record()
    query = build_query_plan(record)[0].query
    payload = _outcome_to_dict(
        ShareholderResearcher(
            StaticSearchClient(), SearchPolicy(delay_seconds=0)
        ).research_person(record)
    )
    payload["ai_overview_log"] = [
        {
            "excel_row": 2,
            "folio_number": "F",
            "primary_holder": record.name,
            "holder_associated": record.name,
            "query_strategy": "Q1",
            "search_query": query,
            "capture_timestamp": "t",
            "overview_status": STATUS_PRESENT,
            "raw_overview_text": LIVE_FALSE_POSITIVE_TEXT,
            "source_titles": ["Privacy Policy"],
            "source_urls": ["https://policies.google.com/privacy?hl=en-IN"],
            "extraction_notes": "",
        }
    ]
    payload["ai_contact_leads"] = [
        {
            "phone": "2147483648",
            "institution": "var",
            "holder_associated": record.name,
            "search_query": query,
            "supporting_evidence_text": LIVE_FALSE_POSITIVE_TEXT,
            "verification_status": "UNVERIFIED",
            "extraction_notes": "",
        }
    ]
    restored = _outcome_from_dict(payload)
    assert restored.ai_contact_leads == []
    assert restored.ai_overview_log[0].overview_status != STATUS_PRESENT


def test_dedupe_log_by_query_not_raw_text():
    first = AiOverviewLogEntry(
        excel_row=2,
        folio_number="",
        primary_holder="A",
        holder_associated="A",
        query_strategy="Q1",
        search_query="same query",
        capture_timestamp="t1",
        overview_status=STATUS_NOT_PRESENT,
    )
    second = AiOverviewLogEntry(
        excel_row=2,
        folio_number="",
        primary_holder="A",
        holder_associated="A",
        query_strategy="Q1",
        search_query="same query",
        capture_timestamp="t2",
        overview_status=STATUS_UNAVAILABLE,
    )
    from src.scraper.ai_overview import dedupe_log_entries, merge_overview_log

    merged = merge_overview_log([first], second)
    assert len(merged) == 1
    assert merged[0].overview_status == STATUS_UNAVAILABLE


def test_playwright_payload_states():
    unavailable = capture_from_playwright_payload(
        {"state": "unavailable", "text": "", "sources": [], "notes": "n/a"}
    )
    assert unavailable.status == STATUS_UNAVAILABLE
    present = capture_from_playwright_payload(
        {
            "state": "present",
            "text": "Jane Example email help@example.org phone +44 20 7946 0958",
            "sources": [{"title": "Example", "url": "https://example.org/p"}],
            "truncated": False,
            "notes": "",
        }
    )
    assert present.status == STATUS_PRESENT


def test_ai_overview_alone_never_strong_match():
    record = _record()
    query = build_query_plan(record)[0].query
    client = StaticSearchClient()
    client.register(query, [], ai_overview=parse_ai_overview_from_html(OVERVIEW_HTML))
    outcome = ShareholderResearcher(client, SearchPolicy(delay_seconds=0)).research_person(record)
    assert outcome.status != ResearchStatus.STRONG_MATCH
    assert outcome.ai_contact_leads
    assert not outcome.strong_findings


def test_truncated_overview_notes():
    capture = parse_ai_overview_from_html(TRUNCATED_HTML)
    assert capture.status == STATUS_PRESENT
    assert "truncated" in capture.extraction_notes.casefold()


def test_empty_overview_container_fails_gracefully():
    html = '<div data-attrid="SGCOP"><span>AI Overview</span></div>'
    capture = parse_ai_overview_from_html(html)
    assert capture.status == STATUS_CAPTURE_FAILED


def test_dry_run_fixtures_summary():
    cases = {
        "unavailable_live": parse_ai_overview_from_html(UNAVAILABLE_HTML),
        "genuine": parse_ai_overview_from_html(OVERVIEW_HTML),
    }
    assert cases["unavailable_live"].status == STATUS_UNAVAILABLE
    assert cases["genuine"].status == STATUS_PRESENT
    assert "nava@usm.my" in cases["genuine"].raw_text


def test_workbook_export_ai_rows(tmp_path):
    source = tmp_path / "in.xlsx"
    destination = tmp_path / "out.xlsx"
    book = Workbook()
    sheet = book.active
    sheet.append(["CL-ID/Folio No.", "Name of the Holder", "Address-1", "Address-2", "Country"])
    sheet.append(["F9", "VISWESWARAN NAVARATNAM", "12 RESEARCH PARK", "PENANG", "Malaysia"])
    book.save(source)
    book.close()

    query = build_query_plan(_record())[0].query
    client = StaticSearchClient()
    client.register(query, [], ai_overview=parse_ai_overview_from_html(OVERVIEW_HTML))
    run_workbook(source, destination, client, SearchPolicy(delay_seconds=0), excel_row=2)

    saved = load_workbook(destination)
    present_rows = [
        row
        for row in saved["AI Overview Log"].iter_rows(min_row=2, values_only=True)
        if row[7] == STATUS_PRESENT
    ]
    assert present_rows
    assert list(saved["AI Overview Log"].iter_rows(values_only=True))[0] == tuple(
        AI_OVERVIEW_LOG_COLUMNS
    )
    saved.close()
