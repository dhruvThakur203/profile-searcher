import asyncio
import random
from datetime import datetime, timedelta, timezone

from openpyxl import Workbook, load_workbook

from src.matching import evaluate_candidate
from src.models import Candidate, MatchStatus, PersonRecord
from src.query_planner import FAMILY_HOLDER_ADDRESS, build_query_plan
from src.scraper.block_cache import GoogleBlockCache
from src.scraper.checkpoint import RowCheckpoint, SqliteCheckpoint
from src.scraper.excel_io import load_input
from src.scraper.google_config import (
    ACCEPT_LANGUAGE,
    CHROME_LAUNCH_ARGS,
    WEBDRIVER_INIT_SCRIPT,
    GoogleBrowserConfig,
)
from src.scraper.google_result_parser import clean_google_results, parse_google_html
from src.scraper.google_state import GooglePageState, detect_google_page
from src.scraper.playwright_google_client import (
    SEARCH_BOX_SELECTORS,
    ChromeLaunchError,
    PageSnapshot,
    PlaywrightGoogleDriver,
    PlaywrightGoogleSearchClient,
    SearchBoxInteractionError,
    SearchBoxNotFound,
    bbox_intersects_viewport,
    first_search_box_selector,
)
from src.scraper.urls import is_valid_external_source_url
from src.scraper.phones import plausible_phone
from src.scraper.researcher import ResearchStatus, ShareholderResearcher
from src.scraper.search_client import SearchPolicy, SearchStatus
from src.scraper.timing import SearchTiming
from src.scraper.urls import (
    canonical_url,
    extract_visible_destination_url,
    resolve_result_url,
    unwrap_google_href,
)
from src.tools.search import SearchResult


SERP_HTML = """
<html><body>
<div id="search">
  <div class="g">
    <a href="/url?q=https://records.example/a.pdf&amp;sa=U"><h3>Record A</h3></a>
    <div class="VwiC3b">Public record snippet</div>
  </div>
  <div class="g">
    <a href="https://records.example/b.html"><h3>Record B</h3></a>
    <div class="IsZvec">Second snippet</div>
  </div>
  <a href="https://www.google.com/search?q=next"><h3>Next</h3></a>
</div>
</body></html>
"""


def test_shareholder_queries_are_generated_and_deduplicated():
    record = PersonRecord(
        record_id="1",
        name="TEST PERSON",
        address1="16708 AMBERWOOD WAY",
        address2="UNIT 4",
        address3="CITYVILLE",
        address4="",
        pincode="560001",
        country="India",
        scrip="Example Ltd",
        joint1="Other Human",
    )
    plans = build_query_plan(record)
    queries = [plan.query for plan in plans]
    strategies = [plan.strategy for plan in plans]

    assert len(queries) == len(set(queries))
    assert "Q1_FULL_NAME_FULL_ADDRESS" in strategies
    assert "Q_NAME_ADDRESS3" in strategies
    assert "Q_NAME_PINCODE" in strategies
    assert "Q_NAME_ADDRESS_FRAGMENT" in strategies
    assert "Q4_NAME_COUNTRY" in strategies
    assert "Q5_NAME_SCRIP" in strategies
    full = next(plan.query for plan in plans if plan.strategy == "Q1_FULL_NAME_FULL_ADDRESS")
    assert "16708 AMBERWOOD WAY" in full
    assert "560001" in full
    # A joint holder is searched alone only in its own address variants;
    # otherwise it appears alongside another holder name.
    for plan in plans:
        if "Other Human" in plan.query:
            if plan.family == FAMILY_HOLDER_ADDRESS:
                assert plan.strategy.endswith("_HOLDER2")
                assert "TEST PERSON" not in plan.query
                continue
            assert plan.strategy.startswith(("Q_JOINT_", "Q_OBIT_JOINT_", "Q_HOLDER_", "Q_ALL_HOLDERS"))
            assert "TEST PERSON" in plan.query
    assert any(plan.query == "TEST PERSON AMBERWOOD WAY" for plan in plans)


def test_google_html_parser_unwraps_results_and_drops_google_links():
    results = parse_google_html(SERP_HTML)

    assert [(item.rank, item.title, item.url, item.snippet) for item in results] == [
        (1, "Record A", "https://records.example/a.pdf", "Public record snippet"),
        (2, "Record B", "https://records.example/b.html", "Second snippet"),
    ]


def test_google_redirect_and_url_normalization():
    assert unwrap_google_href("/url?q=https://records.example/a.pdf&sa=U") == (
        "https://records.example/a.pdf"
    )
    assert unwrap_google_href(
        "https://www.google.com/url?url=https%3A%2F%2Fexample.com%2Fpage&sa=U"
    ) == ("https://example.com/page")
    assert canonical_url("https://records.example/a/?utm_source=google#section") == (
        "https://records.example/a"
    )
    assert canonical_url("https://www.google.com/search?q=person") == ""
    assert canonical_url("https://records.example/file.xlsx?id=42") == (
        "https://records.example/file.xlsx?id=42"
    )


def test_classic_google_url_q_unwrap():
    assert (
        resolve_result_url("https://www.google.com/url?q=https%3A%2F%2Fexample.com%2Fpage")
        == "https://example.com/page"
    )


def test_classic_google_url_url_param_unwrap():
    assert (
        resolve_result_url("/url?url=https%3A%2F%2Frecords.example%2Fb.html&sa=U")
        == "https://records.example/b.html"
    )


def test_goto_href_recovers_visible_destination_from_same_card():
    href = "/goto?url=CAEScgHrOzAVKbRzmUxelGrEkmfyF2Zy4Mc4nQO2lCVL9D8R"
    card_text = (
        "Ravi Cidambi (Age 76) Laguna Hills, California "
        "TruthFinder https://www.truthfinder.com/people-search/ravi-cida "
        "Phone number is (562) 926-1202"
    )
    assert resolve_result_url(href, card_text) == (
        "https://www.truthfinder.com/people-search/ravi-cida"
    )


def test_goto_href_without_visible_destination_does_not_invent_url():
    href = "/goto?url=CAEScgHrOzAVKbRzmUxelGrEkmfyF2Zy4Mc4nQO2lCVL9D8R"
    assert resolve_result_url(href, "Ravi Cidambi Laguna Hills TruthFinder") == ""
    assert extract_visible_destination_url("truthfinder.com people search") == ""


def test_visible_google_destination_is_filtered():
    href = "/goto?url=CAES opaque"
    card_text = "Help https://support.google.com/websearch/answer/181196"
    assert resolve_result_url(href, card_text) == ""


def test_visible_external_destination_survives_canonicalization():
    href = "https://www.google.com/goto?url=CAESopaque"
    card_text = "Directory https://www.usphonebook.com/Ravi-Cidambi?utm_source=google#top"
    assert resolve_result_url(href, card_text) == "https://www.usphonebook.com/Ravi-Cidambi"


def test_multiple_goto_cards_keep_their_own_destination_urls():
    items = [
        {
            "title": "TruthFinder result",
            "url": "/goto?url=CAES111",
            "snippet": "Age 76",
            "card_text": "TruthFinder https://www.truthfinder.com/ravi Age 76",
        },
        {
            "title": "USPhoneBook result",
            "url": "/goto?url=CAES222",
            "snippet": "Also known as",
            "card_text": "USPhoneBook https://www.usphonebook.com/Ravi Also known as",
        },
    ]
    cleaned = clean_google_results(items)
    assert [(item.title, item.url) for item in cleaned] == [
        ("TruthFinder result", "https://www.truthfinder.com/ravi"),
        ("USPhoneBook result", "https://www.usphonebook.com/Ravi"),
    ]


def test_normal_external_href_is_unchanged():
    assert (
        resolve_result_url("https://records.example/b.html", "ignored https://other.example/x")
        == "https://records.example/b.html"
    )


def test_goto_serp_html_recovers_visible_urls():
    html = """
    <html><body><div id="search"><div id="rso">
      <div class="MjjYud">
        <div class="yuRUbf">
          <a href="/goto?url=CAES111"><h3>Ravi Cidambi (Age 76)</h3></a>
          <cite>https://www.truthfinder.com/people-search/ravi</cite>
        </div>
        <div class="VwiC3b">Phone number is (562) 926-1202</div>
      </div>
      <div class="MjjYud">
        <div class="yuRUbf">
          <a href="/goto?url=CAES222"><h3>Ravi Cidambi in Laguna Hills</h3></a>
          <cite>https://www.usphonebook.com/Ravi-Cidambi</cite>
        </div>
        <div class="VwiC3b">Also known as: Cidambi S Ravi</div>
      </div>
    </div></div></body></html>
    """
    results = parse_google_html(html)
    assert [item.url for item in results] == [
        "https://www.truthfinder.com/people-search/ravi",
        "https://www.usphonebook.com/Ravi-Cidambi",
    ]


def test_captcha_page_is_detected_and_not_treated_as_a_serp():
    assert (
        detect_google_page(
            "https://www.google.com/sorry/index",
            "<html>unusual traffic</html>",
        )
        == GooglePageState.CAPTCHA
    )
    assert (
        detect_google_page(
            "https://www.google.com/search?q=person",
            '<html><div id="search">results</div></html>',
        )
        == GooglePageState.SERP
    )
    assert (
        detect_google_page(
            "https://consent.google.com/m",
            "<html><button>Accept all</button></html>",
        )
        == GooglePageState.CONSENT
    )


def test_search_box_selectors_prefer_textarea():
    assert first_search_box_selector(list(SEARCH_BOX_SELECTORS)) == 'textarea[name="q"]'
    assert first_search_box_selector(['input[name="q"]']) == 'input[name="q"]'
    assert first_search_box_selector([]) is None


def test_captcha_sets_cooldown_and_stops_further_google_searches(tmp_path):
    class FakeDriver:
        def __init__(self) -> None:
            self.queries = []
            self.submitted = False

        async def start(self) -> None:
            return None

        async def stop(self) -> None:
            return None

        async def ensure_home(self) -> None:
            return None

        async def submit_query(self, query: str, timing: SearchTiming) -> None:
            self.submitted = True
            self.queries.append(query)

        async def settle_reading(self) -> None:
            return None

        async def snapshot(self) -> PageSnapshot:
            if not self.submitted:
                return PageSnapshot(
                    url="https://www.google.com/",
                    html="<html><textarea name=\"q\"></textarea></html>",
                    title="Google",
                )
            return PageSnapshot(
                url="https://www.google.com/sorry/index",
                html="<html>Our systems have detected unusual traffic</html>",
                title="Sorry",
            )

        async def content(self) -> str:
            return ""

        async def dismiss_consent(self) -> bool:
            return False

        async def evaluate_results(self) -> list[dict]:
            return []

        async def locator_results(self) -> list[dict]:
            return []

        async def save_diagnostics(self, directory: str, snapshot: PageSnapshot) -> None:
            target = tmp_path / "diagnostics"
            target.mkdir(parents=True, exist_ok=True)
            (target / "google_block.html").write_text(snapshot.html, encoding="utf-8")

    now = datetime(2026, 10, 3, tzinfo=timezone.utc)
    cache = GoogleBlockCache(
        tmp_path / "google_block.json",
        cooldown_seconds=1800,
        clock=lambda: now,
    )
    driver = FakeDriver()
    client = PlaywrightGoogleSearchClient(
        GoogleBrowserConfig(
            debug=True,
            diagnostics_dir=str(tmp_path / "diagnostics"),
            timing=SearchTiming.instant(block_cooldown_seconds=1800),
        ),
        driver=driver,
        block_cache=cache,
        sleeper=_async_noop,
    )

    async def both():
        first = await client.search("TEST PERSON 123 TEST STREET")
        second_response = await client.search("TEST PERSON UNIT 4")
        return first, second_response

    blocked, second = asyncio.run(both())

    assert blocked.status == SearchStatus.BLOCKED
    assert blocked.halted is True
    assert second.status == SearchStatus.BLOCKED
    assert driver.queries == ["TEST PERSON 123 TEST STREET"]
    assert cache.is_blocked()
    assert (tmp_path / "diagnostics" / "google_block.html").exists()
    assert "solve" not in blocked.message.casefold()
    assert "duckduckgo" not in blocked.message.casefold()
    assert "bing" not in blocked.message.casefold()


def test_html_fallback_is_used_when_javascript_extraction_fails(tmp_path):
    class FakeDriver:
        async def start(self) -> None:
            return None

        async def stop(self) -> None:
            return None

        async def ensure_home(self) -> None:
            return None

        async def submit_query(self, query: str, timing: SearchTiming) -> None:
            return None

        async def settle_reading(self) -> None:
            return None

        async def snapshot(self) -> PageSnapshot:
            return PageSnapshot(
                url="https://www.google.com/search?q=test",
                html=SERP_HTML,
                title="results",
            )

        async def content(self) -> str:
            return SERP_HTML

        async def dismiss_consent(self) -> bool:
            return False

        async def evaluate_results(self) -> list[dict]:
            raise RuntimeError("javascript failed")

        async def locator_results(self) -> list[dict]:
            return []

        async def save_diagnostics(self, directory: str, snapshot: PageSnapshot) -> None:
            return None

    client = PlaywrightGoogleSearchClient(
        GoogleBrowserConfig(timing=SearchTiming.instant()),
        driver=FakeDriver(),
        block_cache=GoogleBlockCache(
            tmp_path / "block.json",
            cooldown_seconds=1,
            clock=lambda: datetime(2026, 1, 1, tzinfo=timezone.utc),
        ),
        sleeper=_async_noop,
    )
    # Point the cache at a missing path by replacing after init if needed.
    response = asyncio.run(client.search("TEST PERSON"))

    assert response.status == SearchStatus.OK
    assert [item.url for item in response.results] == [
        "https://records.example/a.pdf",
        "https://records.example/b.html",
    ]


def test_missing_search_box_is_a_controlled_error(tmp_path):
    class FakeDriver:
        async def start(self) -> None:
            return None

        async def stop(self) -> None:
            return None

        async def ensure_home(self) -> None:
            return None

        async def submit_query(self, query: str, timing: SearchTiming) -> None:
            raise SearchBoxNotFound("Google search box was not found.")

        async def settle_reading(self) -> None:
            return None

        async def snapshot(self) -> PageSnapshot:
            return PageSnapshot(url="https://www.google.com/", html="<html></html>", title="Google")

        async def content(self) -> str:
            return ""

        async def dismiss_consent(self) -> bool:
            return False

        async def evaluate_results(self) -> list[dict]:
            return []

        async def locator_results(self) -> list[dict]:
            return []

        async def save_diagnostics(self, directory: str, snapshot: PageSnapshot) -> None:
            return None

    client = PlaywrightGoogleSearchClient(
        GoogleBrowserConfig(timing=SearchTiming.instant()),
        driver=FakeDriver(),
        block_cache=GoogleBlockCache(tmp_path / "block.json", cooldown_seconds=1),
        sleeper=_async_noop,
    )
    response = asyncio.run(client.search("TEST PERSON"))
    assert response.status == SearchStatus.ERROR
    assert response.results == []


def test_phone_validation_rejects_ids_dates_and_pins():
    assert plausible_phone("9999999999") == "9999999999"
    assert plausible_phone("+1 555 123456") == "+1 555 123456"
    assert plausible_phone("01/02/2024") is None
    assert plausible_phone("12345678901234") is None
    assert plausible_phone("560001") is None
    assert plausible_phone("9999999999", blocked_digits={"9999999999"}) is None


def test_exact_name_and_address_threshold_still_hold():
    person = PersonRecord(
        record_id="1",
        name="Dhruv Singh Thakur",
        address1="16708 AMBERWOOD WAY",
    )
    assert evaluate_candidate(
        person,
        Candidate(name="Dhruv Thakur", address1="16708 AMBERWOOD WAY"),
    ).status == MatchStatus.REJECTED
    assert evaluate_candidate(
        person,
        Candidate(name="Dhruv Singh Thakur", address1="16708 AMBERWOOD WAY"),
    ).status == MatchStatus.STRONG
    weak = evaluate_candidate(
        person,
        Candidate(name="Dhruv Singh Thakur", address1="16708 AMBERWOOD DR"),
    )
    assert weak.address1_similarity < 95
    assert weak.status == MatchStatus.CANDIDATE
    address3 = evaluate_candidate(
        PersonRecord(record_id="1", name="Dhruv Singh Thakur", address3="16708 AMBERWOOD WAY"),
        Candidate(name="Dhruv Singh Thakur", address3="16708 AMBERWOOD WAY"),
    )
    assert address3.status == MatchStatus.STRONG


def test_checkpoint_skips_completed_strategies_and_finished_rows(tmp_path):
    store = SqliteCheckpoint(tmp_path / "checkpoint.sqlite")
    record = PersonRecord(
        record_id="9",
        name="TEST PERSON",
        address1="123 TEST STREET",
        address2="UNIT 4",
    )
    plans = build_query_plan(record)
    q1 = plans[0]
    store.save(
        record.record_id,
        RowCheckpoint(
            finished=False,
            completed_strategies=[q1.strategy],
            findings=[
                {
                    "url": "https://source.example/a",
                    "title": "A",
                    "mobiles": [],
                    "emails": ["a@example.com"],
                    "address1_similarity": 100.0,
                    "address2_similarity": None,
                    "address3_similarity": None,
                    "address4_similarity": None,
                    "strategies": [q1.strategy],
                }
            ],
        ),
    )
    client_calls = []

    class Recording:
        def search(self, query: str):
            client_calls.append(query)
            from src.scraper.search_client import SearchResponse

            if query == plans[1].query:
                return SearchResponse(
                    results=[
                        SearchResult(
                            title="B",
                            url="https://source.example/b",
                            snippet="TEST PERSON 123 TEST STREET",
                        )
                    ]
                )
            return SearchResponse(results=[])

    outcome = ShareholderResearcher(
        Recording(),
        SearchPolicy(delay_seconds=0),
        checkpoint=store,
    ).research_person(record)

    assert q1.query not in client_calls
    assert "https://source.example/a" in outcome.source_links
    assert "https://source.example/b" in outcome.source_links
    assert outcome.emails == ["a@example.com"]

    finished = store.load(record.record_id)
    assert finished is not None and finished.finished
    ShareholderResearcher(
        _Exploding(),
        SearchPolicy(delay_seconds=0),
        checkpoint=store,
    ).research_person(record)


def test_shareholder_workbook_headers_and_existing_output_columns(tmp_path):
    source = tmp_path / "holders.xlsx"
    book = Workbook()
    sheet = book.active
    sheet.append(
        [
            "CL-ID/Folio No.",
            "Name of the Holder",
            "Second Holder (Joint 1)",
            "Address-1",
            "Address-2",
            "Address-3",
            "Pincode",
            " Country",
            "Mobile",
        ]
    )
    sheet.append(
        [
            "FOLIO1",
            "TEST PERSON",
            "Other Human",
            "123 TEST STREET",
            "UNIT 4",
            "",
            "560001",
            "India",
            "old-value",
        ]
    )
    book.save(source)
    book.close()

    _, headers, rows = load_input(source)
    assert rows[0].record.name == "TEST PERSON"
    assert rows[0].record.address1 == "123 TEST STREET"
    assert rows[0].record.address2 == "UNIT 4"
    assert rows[0].record.pincode == "560001"
    assert rows[0].record.country == "India"
    assert rows[0].record.joint1 == "Other Human"
    assert rows[0].excel_row == 2

    from src.scraper.excel_io import write_output
    from src.scraper.researcher import PersonResearch

    outcome = PersonResearch(
        record_id=rows[0].record.record_id,
        input_name="TEST PERSON",
        status=ResearchStatus.STRONG_MATCH,
        matched_name="TEST PERSON",
        mobiles=["9999999999"],
        source_links=["https://source.example/a"],
    )
    destination = tmp_path / "out.xlsx"
    write_output(destination, "Sheet", headers, [(rows[0], outcome)])
    saved = load_workbook(destination)
    saved_headers = [cell.value for cell in saved.active[1]]
    assert saved_headers.count("Mobile") == 1
    mobile_index = saved_headers.index("Mobile")
    assert saved.active[2][mobile_index].value == "9999999999"
    assert saved.active[2][0].value == "FOLIO1"
    saved.close()


def test_normal_google_workflow_does_not_require_brave_api_key(monkeypatch):
    monkeypatch.delenv("BRAVE_SEARCH_API_KEY", raising=False)
    from src.scraper.__main__ import build_parser, create_google_client

    args = build_parser().parse_args(["holders.xlsx", "out.xlsx"])
    assert args.output == "out.xlsx"
    client = create_google_client(
        GoogleBrowserConfig(
            block_path="data/unused-block.json",
            timing=SearchTiming.instant(),
        )
    )
    assert client.manages_browser is True
    client.close()


def test_search_timing_defaults_are_conservative():
    timing = SearchTiming()
    assert (timing.pre_search_min, timing.pre_search_max) == (0.5, 2.5)
    assert (timing.typing_min_ms, timing.typing_max_ms) == (80, 250)
    assert timing.typing_pause_probability == 0.08
    assert (timing.typing_pause_min, timing.typing_pause_max) == (0.15, 0.45)
    assert (timing.post_submit_min, timing.post_submit_max) == (2.0, 8.0)
    assert (timing.between_searches_min, timing.between_searches_max) == (15.0, 45.0)
    assert timing.cooldown_after_searches == 5
    assert (timing.cooldown_min, timing.cooldown_max) == (60.0, 120.0)
    assert (timing.reading_min, timing.reading_max) == (3.0, 10.0)
    assert timing.block_cooldown_seconds == 30 * 60
    config = GoogleBrowserConfig()
    assert config.home_url == "https://www.google.com"
    assert config.headless is False
    assert config.locale == "en-IN"
    assert config.accept_language == ACCEPT_LANGUAGE
    assert (config.viewport_width, config.viewport_height) == (1366, 900)
    assert config.channel == "chrome"
    launch = config.persistent_launch_kwargs()
    assert launch["channel"] == "chrome"
    assert launch["args"] == list(CHROME_LAUNCH_ARGS)
    assert launch["extra_http_headers"]["Accept-Language"] == ACCEPT_LANGUAGE
    assert "Object.defineProperty(navigator, 'webdriver'" in WEBDRIVER_INIT_SCRIPT
    assert ChromeLaunchError is not None


class _Exploding:
    def search(self, query: str):
        raise AssertionError(query)


async def _async_noop(seconds: float) -> None:
    return None


def test_block_cache_roundtrip(tmp_path):
    start = datetime(2026, 10, 3, tzinfo=timezone.utc)
    current = {"now": start}
    cache = GoogleBlockCache(tmp_path / "google_block.json", cooldown_seconds=1800, clock=lambda: current["now"])
    assert cache.is_blocked() is False
    cache.record_block("captcha")
    assert cache.is_blocked() is True
    assert cache.reason() == "captcha"
    current["now"] = start + timedelta(seconds=1801)
    assert cache.is_blocked() is False


def test_timing_samples_stay_within_configured_ranges():
    timing = SearchTiming()
    rng = random.Random(7)
    for _ in range(40):
        assert 0.5 <= timing.sample_pre_search(rng) <= 2.5
        assert 2.0 <= timing.sample_post_submit(rng) <= 8.0
        assert 15.0 <= timing.sample_between_searches(rng) <= 45.0
        assert 60.0 <= timing.sample_cooldown(rng) <= 120.0
        assert 3.0 <= timing.sample_reading(rng) <= 10.0
        assert 80 <= timing.sample_typing_delay_ms(rng) <= 250


def test_typing_delays_are_randomized_per_character():
    timing = SearchTiming()
    rng = random.Random(11)
    delays = [timing.sample_typing_delay_ms(rng) for _ in range(20)]
    assert min(delays) >= 80
    assert max(delays) <= 250
    assert len(set(delays)) > 1


def test_typing_pause_probability_and_range():
    timing = SearchTiming(
        typing_pause_probability=1.0,
        typing_pause_min=0.15,
        typing_pause_max=0.45,
    )
    rng = random.Random(3)
    pauses = [timing.sample_typing_pause(rng) for _ in range(20)]
    assert all(pause is not None and 0.15 <= pause <= 0.45 for pause in pauses)

    never = SearchTiming(typing_pause_probability=0.0)
    assert all(never.sample_typing_pause(random.Random(i)) is None for i in range(20))


def test_chrome_launch_config_has_no_chromium_fallback(tmp_path):
    config = GoogleBrowserConfig(profile_dir=str(tmp_path / "browser_profile"))
    kwargs = config.persistent_launch_kwargs()
    assert kwargs["channel"] == "chrome"
    assert kwargs["user_data_dir"] == str(tmp_path / "browser_profile")
    assert "--disable-blink-features=AutomationControlled" in kwargs["args"]
    assert "--disable-infobars" in kwargs["args"]
    assert "--no-sandbox" in kwargs["args"]
    assert kwargs["extra_http_headers"]["Accept-Language"] == "en-IN,en;q=0.9"
    assert "chromium" not in kwargs
    assert "executable_path" not in kwargs


def test_one_page_is_reused_and_cooldown_counts_only_successful_searches(tmp_path):
    class TrackingDriver:
        def __init__(self) -> None:
            self.start_count = 0
            self.queries = []
            self.settle_count = 0

        async def start(self) -> None:
            self.start_count += 1

        async def stop(self) -> None:
            return None

        async def ensure_home(self) -> None:
            return None

        async def submit_query(self, query: str, timing: SearchTiming) -> None:
            self.queries.append(query)

        async def settle_reading(self) -> None:
            self.settle_count += 1

        async def snapshot(self) -> PageSnapshot:
            return PageSnapshot(
                url="https://www.google.com/search?q=test",
                html=SERP_HTML,
                title="results",
            )

        async def content(self) -> str:
            return SERP_HTML

        async def dismiss_consent(self) -> bool:
            return False

        async def evaluate_results(self) -> list[dict]:
            return [{"title": "TEST PERSON", "url": "https://records.example/a.pdf", "snippet": ""}]

        async def locator_results(self) -> list[dict]:
            return []

        async def save_diagnostics(self, directory: str, snapshot: PageSnapshot) -> None:
            return None

    sleeps = []

    async def record_sleep(seconds: float) -> None:
        sleeps.append(seconds)

    timing = SearchTiming.instant(
        cooldown_after_searches=5,
        cooldown_min=60.0,
        cooldown_max=120.0,
        between_searches_min=15.0,
        between_searches_max=45.0,
    )
    driver = TrackingDriver()
    client = PlaywrightGoogleSearchClient(
        GoogleBrowserConfig(
            profile_dir=str(tmp_path / "shared-profile"),
            timing=timing,
        ),
        driver=driver,
        block_cache=GoogleBlockCache(tmp_path / "block.json", cooldown_seconds=1),
        sleeper=record_sleep,
        rng=random.Random(21),
    )

    async def run_six():
        responses = []
        for index in range(6):
            responses.append(await client.search(f"QUERY {index}"))
        return responses

    responses = asyncio.run(run_six())
    assert all(item.status == SearchStatus.OK for item in responses)
    assert driver.start_count == 1
    assert len(driver.queries) == 6
    assert driver.settle_count == 6
    between = [value for value in sleeps if 15.0 <= value <= 45.0]
    assert len(between) == 5
    cooldowns = [value for value in sleeps if 60.0 <= value <= 120.0]
    assert len(cooldowns) == 1
    assert 60.0 <= cooldowns[0] <= 120.0


def test_blocked_search_does_not_count_toward_five_search_cooldown(tmp_path):
    class BlockThenOkDriver:
        def __init__(self) -> None:
            self.queries = []

        async def start(self) -> None:
            return None

        async def stop(self) -> None:
            return None

        async def ensure_home(self) -> None:
            return None

        async def submit_query(self, query: str, timing: SearchTiming) -> None:
            self.queries.append(query)

        async def settle_reading(self) -> None:
            return None

        async def snapshot(self) -> PageSnapshot:
            if len(self.queries) == 0:
                return PageSnapshot(
                    url="https://www.google.com/",
                    html='<html><textarea name="q"></textarea></html>',
                    title="Google",
                )
            if len(self.queries) == 1:
                return PageSnapshot(
                    url="https://www.google.com/sorry/index",
                    html="<html>unusual traffic</html>",
                    title="Sorry",
                )
            return PageSnapshot(
                url="https://www.google.com/search?q=ok",
                html=SERP_HTML,
                title="results",
            )

        async def content(self) -> str:
            return SERP_HTML

        async def dismiss_consent(self) -> bool:
            return False

        async def evaluate_results(self) -> list[dict]:
            return [{"title": "TEST PERSON", "url": "https://records.example/a.pdf", "snippet": ""}]

        async def locator_results(self) -> list[dict]:
            return []

        async def save_diagnostics(self, directory: str, snapshot: PageSnapshot) -> None:
            return None

    sleeps = []

    async def record_sleep(seconds: float) -> None:
        sleeps.append(seconds)

    # First search blocks; completed_searches stays 0, so no 5-search cooldown.
    client = PlaywrightGoogleSearchClient(
        GoogleBrowserConfig(timing=SearchTiming.instant(cooldown_min=60, cooldown_max=120)),
        driver=BlockThenOkDriver(),
        block_cache=GoogleBlockCache(tmp_path / "block.json", cooldown_seconds=1),
        sleeper=record_sleep,
        rng=random.Random(5),
    )
    blocked = asyncio.run(client.search("BLOCKED QUERY"))
    assert blocked.status == SearchStatus.BLOCKED
    assert not any(60.0 <= value <= 120.0 for value in sleeps)


def test_bbox_intersects_viewport_helper():
    assert bbox_intersects_viewport(
        {"x": 10, "y": 10, "width": 100, "height": 20}, 1366, 900
    )
    assert not bbox_intersects_viewport(
        {"x": 10, "y": 2000, "width": 100, "height": 20}, 1366, 900
    )
    assert not bbox_intersects_viewport(None, 1366, 900)


def test_resource_and_invalid_source_urls_are_rejected():
    assert is_valid_external_source_url("https://records.example/a")
    assert is_valid_external_source_url("http://records.example/a")
    assert not is_valid_external_source_url("http://www.w3.org/2000/svg")
    assert not is_valid_external_source_url("http://www.w3.org/2000/svg\\")
    assert not is_valid_external_source_url("https://www.google.com/search?q=x")
    assert not is_valid_external_source_url("https://www.gstatic.com/foo.png")
    assert not is_valid_external_source_url("https://googleusercontent.com/x")
    assert not is_valid_external_source_url("data:text/plain,hi")
    assert not is_valid_external_source_url("blob:https://example.com/1")
    assert not is_valid_external_source_url("javascript:alert(1)")
    assert canonical_url("http://www.w3.org/2000/svg\\") == ""
    assert (
        resolve_result_url(
            "/goto?url=CAES",
            "icon http://www.w3.org/2000/svg\\ and no real destination",
        )
        == ""
    )


class _FakeBox:
    def __init__(self, *, visible=True, bbox=None, fail_clicks=0):
        self.visible = visible
        self.bbox = bbox or {"x": 20, "y": 20, "width": 200, "height": 30}
        self.fail_clicks = fail_clicks
        self.click_count = 0
        self.filled = None

    async def is_visible(self):
        return self.visible

    async def scroll_into_view_if_needed(self, timeout=5000):
        self.visible = True
        if self.bbox["y"] > 900:
            self.bbox = {**self.bbox, "y": 40}

    async def bounding_box(self):
        return dict(self.bbox)

    async def click(self, timeout=5000):
        self.click_count += 1
        if self.fail_clicks > 0:
            self.fail_clicks -= 1
            raise RuntimeError("ElementHandle.click: element is outside of the viewport")

    async def fill(self, value):
        self.filled = value


class _FakePage:
    def __init__(self, boxes_by_selector):
        self.boxes_by_selector = boxes_by_selector
        self.url = "https://www.google.com/search?q=prior"
        self.viewport_size = {"width": 1366, "height": 900}
        self.typed = []
        self.pressed = []
        self.home_loads = 0

    async def title(self):
        return "Google Search"

    async def query_selector_all(self, selector):
        return list(self.boxes_by_selector.get(selector, []))

    async def evaluate(self, script):
        return None

    async def goto(self, url, wait_until="domcontentloaded"):
        self.home_loads += 1
        self.url = url

    async def wait_for_selector(self, selector, timeout=15000):
        return None


async def _async_noop_sleep(seconds: float) -> None:
    return None


def test_search_box_visible_click_succeeds():
    box = _FakeBox()
    driver = PlaywrightGoogleDriver(GoogleBrowserConfig(), sleeper=_async_noop_sleep)
    driver._page = _FakePage({'textarea[name="q"]': [box]})
    prepared, selector = asyncio.run(driver._find_and_prepare_search_box("QUERY"))
    assert selector == 'textarea[name="q"]'
    assert prepared is box
    assert box.click_count == 1


def test_search_box_outside_viewport_scrolls_then_clicks():
    box = _FakeBox(bbox={"x": 20, "y": 2000, "width": 200, "height": 30})
    driver = PlaywrightGoogleDriver(GoogleBrowserConfig(), sleeper=_async_noop_sleep)
    driver._page = _FakePage({'textarea[name="q"]': [box]})
    prepared, _ = asyncio.run(driver._find_and_prepare_search_box("QUERY"))
    assert prepared is box
    assert box.bbox["y"] < 900
    assert box.click_count == 1


def test_search_box_stale_element_is_relocated():
    first = _FakeBox(fail_clicks=1)
    second = _FakeBox()
    page = _FakePage({'textarea[name="q"]': [first]})

    async def swapping_query_selector_all(selector):
        if first.click_count >= 1:
            return [second]
        return [first]

    page.query_selector_all = swapping_query_selector_all
    driver = PlaywrightGoogleDriver(GoogleBrowserConfig(), sleeper=_async_noop_sleep)
    driver._page = page
    prepared, _ = asyncio.run(driver._find_and_prepare_search_box("QUERY"))
    assert prepared is second
    assert first.click_count == 1
    assert second.click_count == 1


def test_search_box_bounded_retry_then_controlled_error():
    box = _FakeBox(fail_clicks=10)
    driver = PlaywrightGoogleDriver(GoogleBrowserConfig(), sleeper=_async_noop_sleep)
    driver._page = _FakePage({'textarea[name="q"]': [box]})
    try:
        asyncio.run(driver._find_and_prepare_search_box("QUERY"))
        raised = False
    except SearchBoxInteractionError:
        raised = True
    assert raised
    assert box.click_count == 3
